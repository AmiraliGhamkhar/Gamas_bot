"""Pure Office export modules retained after the bot UI retirement.

These tests keep the standalone builders honest: source wording stays intact,
model-proposed content must be grounded in that source, and every export is a
real package a reader can open. The Telegram UI and job lane are no longer part
of the product surface.
"""

from __future__ import annotations

import io
import json as json_lib
import tempfile
import unittest
import zipfile
from pathlib import Path

import docx as docx_library
import openpyxl
from pptx import Presentation

from gamas_bot.office import (
    OfficeError,
    build_office_document,
    build_office_pptx,
    build_office_xlsx,
    fallback_slides,
    fallback_table,
    office_filename,
    office_title_from_filename,
    read_text_file,
)
from gamas_bot.office_ai import _chunk_long_text, infer_table, plan_deck
from gamas_bot.mdblocks import heading_level, parse_blocks
from tests.support import (
    docx_text,
    make_settings,
)

SAMPLE = (
    "# گزارش عملکرد\n"
    "## بخش اول\n"
    "این متن عیناً در سند می‌آید.\n"
    "- مورد اول\n"
    "- مورد دوم\n"
    "## بخش دوم\n"
    "نام\tسن\tشهر\n"
    "علی\t۲۰\tتهران\n"
    "سارا\t25\tشیراز\n"
)


# ---------------------------------------------------------------------------
# The parser
# ---------------------------------------------------------------------------


class BlockParserTests(unittest.TestCase):
    def test_headings_bullets_and_paragraphs(self):
        blocks = parse_blocks(SAMPLE)
        kinds = [block.kind for block in blocks]
        self.assertEqual(kinds[:4], ["heading", "heading", "paragraph", "bullet"])
        self.assertEqual(blocks[0].text, "گزارش عملکرد")
        self.assertEqual(blocks[0].level, 1)
        self.assertEqual(blocks[1].level, 2)
        self.assertEqual(blocks[3].text, "مورد اول")

    def test_tab_separated_rows_become_a_table(self):
        blocks = parse_blocks(SAMPLE)
        tables = [block for block in blocks if block.is_table]
        self.assertEqual(len(tables), 1)
        self.assertEqual(
            tables[0].rows,
            (
                ("نام", "سن", "شهر"),
                ("علی", "۲۰", "تهران"),
                ("سارا", "25", "شیراز"),
            ),
        )

    def test_pipe_table_rows_become_a_table(self):
        blocks = parse_blocks("| نام | سن |\n| --- | --- |\n| علی | ۲۰ |\n")
        tables = [block for block in blocks if block.is_table]
        self.assertEqual(len(tables), 1)
        self.assertEqual(tables[0].rows, (("نام", "سن"), ("علی", "۲۰")))

    def test_comma_rows_become_a_table(self):
        blocks = parse_blocks("نام,سن,شهر\nعلی,20,تهران\n")
        tables = [block for block in blocks if block.is_table]
        self.assertEqual(len(tables), 1)
        self.assertEqual(tables[0].rows[0], ("نام", "سن", "شهر"))

    def test_prose_with_commas_is_not_a_table(self):
        """Persian prose is full of commas; a two-column guess would mangle it."""
        blocks = parse_blocks("سلام، خوبی؟\nبله، ممنون.\n")
        self.assertEqual([block.kind for block in blocks], ["paragraph", "paragraph"])
        self.assertFalse(any(block.is_table for block in blocks))

    def test_every_line_lands_in_exactly_one_block(self):
        blocks = parse_blocks(SAMPLE)
        self.assertEqual(len(blocks), 7)

    def test_heading_level_is_clamped(self):
        self.assertEqual(heading_level(1), 1)
        self.assertEqual(heading_level(6), 3)
        # The lecture pipeline's convention: its top level is ``##``.
        self.assertEqual(heading_level(2, offset=-1), 1)
        self.assertEqual(heading_level(3, offset=-1), 2)


# ---------------------------------------------------------------------------
# The document model
# ---------------------------------------------------------------------------


class OfficeDocumentTests(unittest.TestCase):
    def test_title_comes_from_the_first_heading(self):
        document = build_office_document(SAMPLE)
        self.assertEqual(document.title, "گزارش عملکرد")

    def test_title_falls_back_to_the_first_line(self):
        document = build_office_document("فهرست هزینه‌ها\nمورد یک\nمورد دو\n")
        self.assertEqual(document.title, "فهرست هزینه‌ها")

    def test_tables_are_extracted(self):
        document = build_office_document(SAMPLE)
        self.assertEqual(len(document.tables), 1)
        self.assertEqual(document.tables[0].headers, ("نام", "سن", "شهر"))

    def test_fallback_slides_follow_the_users_headings(self):
        document = build_office_document(SAMPLE)
        slides = fallback_slides(document)
        self.assertEqual([slide.title for slide in slides], ["بخش اول", "بخش دوم"])
        self.assertIn("مورد اول", slides[0].bullets)

    def test_a_heading_without_a_body_still_becomes_a_slide(self):
        document = build_office_document("# عنوان\n\n## بخش خالی\n## بخش پر\nمتن\n")
        slides = fallback_slides(document)
        self.assertEqual([slide.title for slide in slides], ["بخش خالی", "بخش پر"])
        self.assertEqual(slides[0].bullets, ())

    def test_fallback_slides_without_headings(self):
        document = build_office_document("سطر یک\nسطر دو\n")
        slides = fallback_slides(document)
        self.assertEqual(len(slides), 1)
        # The first line titles the slide instead of repeating itself.
        self.assertEqual(slides[0].title, "سطر یک")
        self.assertEqual(slides[0].bullets, ("سطر دو",))

    def test_a_long_text_without_headings_is_cut_into_slides(self):
        lines = [f"سطر {index}" for index in range(1, 14)]
        document = build_office_document("\n".join(lines))
        slides = fallback_slides(document)
        self.assertEqual(len(slides), 2)
        self.assertEqual(slides[0].title, "سطر 1")
        self.assertEqual(len(slides[0].bullets), 6)
        self.assertEqual(slides[1].title, "بخش ۲")
        self.assertEqual(len(slides[1].bullets), 6)

    def test_fallback_table_keeps_every_line(self):
        document = build_office_document("سطر یک\nسطر دو\n")
        table = fallback_table(document)
        self.assertEqual(table.headers, ("متن",))
        self.assertEqual(table.rows, (("سطر یک",), ("سطر دو",)))

    def test_filename_helpers(self):
        self.assertEqual(office_filename("گزارش", "word", "GMS-7"), "گزارش - GMS-7.docx")
        self.assertEqual(office_filename("گزارش", "excel", "GMS-7"), "گزارش - GMS-7.xlsx")
        self.assertEqual(office_filename("گزارش", "ppt", "GMS-7"), "گزارش - GMS-7.pptx")
        self.assertEqual(office_title_from_filename("گزارش_ماهانه.txt"), "گزارش ماهانه")

    def test_text_file_is_read_in_any_encoding(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "a.txt"
            # cp1256 is the Persian Windows encoding; a file written in it
            # must still be read as the text it holds.
            path.write_bytes("گزارش".encode("cp1256"))
            self.assertEqual(read_text_file(path, max_bytes=10_000), "گزارش")

    def test_oversize_text_file_is_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "a.txt"
            path.write_text("x" * 100, encoding="utf-8")
            with self.assertRaises(OfficeError):
                read_text_file(path, max_bytes=10)


# ---------------------------------------------------------------------------
# The three documents
# ---------------------------------------------------------------------------


class WordDocumentTests(unittest.TestCase):
    def test_the_pasted_table_becomes_a_real_word_table(self):
        from gamas_bot.docx_export import (
            DocumentMeta,
            build_plain_docx,
            resolve_design,
            resolve_fonts,
        )
        from datetime import datetime

        document = build_office_document(SAMPLE)
        payload = build_plain_docx(
            document.title,
            document.text,
            fonts=resolve_fonts(),
            design=resolve_design(),
            meta=DocumentMeta(reference="GMS-1", created_at=datetime.now()),
            mode_label="سند آزاد — Word",
            heading_offset=0,
        )
        self.assertTrue(payload.startswith(b"PK"))
        produced = docx_library.Document(io.BytesIO(payload))
        self.assertEqual(len(produced.tables), 1)
        rows = [[cell.text for cell in row.cells] for row in produced.tables[0].rows]
        self.assertEqual(rows[0], ["نام", "سن", "شهر"])
        self.assertEqual(rows[1], ["علی", "۲۰", "تهران"])
        # The user's text is in the document, verbatim.
        text = docx_text(payload)
        self.assertIn("این متن عیناً در سند می‌آید.", text)
        self.assertIn("سند آزاد — Word", text)

    def test_the_document_is_right_to_left(self):
        from gamas_bot.docx_export import (
            DocumentMeta,
            build_plain_docx,
            resolve_design,
            resolve_fonts,
        )
        from datetime import datetime

        document = build_office_document(SAMPLE)
        payload = build_plain_docx(
            document.title,
            document.text,
            fonts=resolve_fonts(),
            design=resolve_design(),
            meta=DocumentMeta(reference="GMS-1", created_at=datetime.now()),
            heading_offset=0,
        )
        xml = zipfile.ZipFile(io.BytesIO(payload)).read("word/document.xml").decode("utf-8")
        self.assertIn("w:bidi", xml)
        self.assertIn("<w:rtl", xml)


class ExcelDocumentTests(unittest.TestCase):
    def test_the_sheet_is_right_to_left_with_a_frozen_header(self):
        document = build_office_document(SAMPLE)
        payload = build_office_xlsx(document, list(document.tables))
        self.assertTrue(payload.startswith(b"PK"))
        workbook = openpyxl.load_workbook(io.BytesIO(payload))
        sheet = workbook.active
        self.assertTrue(sheet.sheet_view.rightToLeft)
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertEqual(sheet["A1"].value, "نام")
        self.assertEqual(sheet["C2"].value, "تهران")

    def test_ascii_numbers_become_numbers_and_persian_digits_stay_text(self):
        """A numeric column must be summable, but «۲۰» must not turn into 20."""
        document = build_office_document(SAMPLE)
        payload = build_office_xlsx(document, list(document.tables))
        workbook = openpyxl.load_workbook(io.BytesIO(payload))
        sheet = workbook.active
        self.assertEqual(sheet["B2"].value, "۲۰")
        self.assertEqual(sheet["B3"].value, 25)

    def test_zero_padded_identifiers_stay_text(self):
        """A phone number is not a number: coercion would delete the leading 0."""
        document = build_office_document(
            "کد\tمقدار\nتلفن\t09123456789\nکد\t007\nعدد\t1.5\n"
        )
        payload = build_office_xlsx(document, list(document.tables))
        sheet = openpyxl.load_workbook(io.BytesIO(payload)).active
        self.assertEqual(sheet["B2"].value, "09123456789")
        self.assertEqual(sheet["B3"].value, "007")
        self.assertEqual(sheet["B4"].value, 1.5)

    def test_a_text_without_a_table_still_produces_a_sheet(self):
        document = build_office_document("سطر یک\nسطر دو\n")
        payload = build_office_xlsx(document, list(document.tables))
        workbook = openpyxl.load_workbook(io.BytesIO(payload))
        sheet = workbook.active
        self.assertEqual(sheet["A1"].value, "متن")
        self.assertEqual(sheet["A2"].value, "سطر یک")

    def test_the_sheet_name_is_the_title(self):
        document = build_office_document(SAMPLE)
        payload = build_office_xlsx(document, list(document.tables))
        workbook = openpyxl.load_workbook(io.BytesIO(payload))
        self.assertEqual(workbook.sheetnames, ["گزارش عملکرد"])


class PowerPointDocumentTests(unittest.TestCase):
    def test_the_deck_has_a_title_an_agenda_the_sections_and_a_closing(self):
        document = build_office_document(SAMPLE)
        slides = fallback_slides(document)
        payload = build_office_pptx(document, slides)
        self.assertTrue(payload.startswith(b"PK"))
        presentation = Presentation(io.BytesIO(payload))
        # title + agenda + one slide per section + closing
        self.assertEqual(len(presentation.slides), 5)

    def test_every_paragraph_is_right_to_left(self):
        document = build_office_document(SAMPLE)
        slides = fallback_slides(document)
        payload = build_office_pptx(document, slides)
        presentation = Presentation(io.BytesIO(payload))
        for slide in presentation.slides:
            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue
                for paragraph in shape.text_frame.paragraphs:
                    if not paragraph.text.strip():
                        continue
                    p_pr = paragraph._p.find(
                        "{http://schemas.openxmlformats.org/drawingml/2006/main}pPr"
                    )
                    self.assertIsNotNone(p_pr, paragraph.text)
                    self.assertEqual(p_pr.get("rtl"), "1", paragraph.text)
                    self.assertIn(p_pr.get("algn"), {"r", "ctr"}, paragraph.text)

    def test_the_slides_carry_the_users_own_lines(self):
        document = build_office_document(SAMPLE)
        slides = fallback_slides(document)
        payload = build_office_pptx(document, slides)
        presentation = Presentation(io.BytesIO(payload))
        texts = []
        for slide in presentation.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    texts.append(shape.text_frame.text)
        joined = "\n".join(texts)
        self.assertIn("این متن عیناً در سند می‌آید.", joined)
        self.assertIn("گزارش عملکرد", joined)

    def test_the_section_text_and_table_rows_are_kept_in_the_notes(self):
        document = build_office_document(SAMPLE)
        slides = fallback_slides(document)
        payload = build_office_pptx(document, slides)
        presentation = Presentation(io.BytesIO(payload))
        notes = [
            slide.notes_slide.notes_text_frame.text
            for slide in presentation.slides
            if slide.has_notes_slide
        ]
        self.assertTrue(any("این متن عیناً در سند می‌آید." in note for note in notes))
        self.assertTrue(any("نام | سن | شهر" in note for note in notes))
        self.assertTrue(any("علی | ۲۰ | تهران" in note for note in notes))


# ---------------------------------------------------------------------------
# The model layer
# ---------------------------------------------------------------------------


class _FakeLLMSession:
    """Records the request and answers with a canned payload.

    The LLM client posts with ``async with session.post(...)``, so ``post``
    is a plain method that returns an async context manager.
    """

    def __init__(self, payload: dict):
        self.payload = payload
        self.requests: list[dict] = []

    def post(self, url, headers=None, json=None, timeout=None):  # noqa: A002
        self.requests.append({"url": url, "json": json})
        return _FakeResponse(
            {"choices": [{"message": {"content": json_lib.dumps(self.payload)}}]}
        )


class _FakeResponse:
    def __init__(self, payload: dict):
        self.payload = payload
        self.status = 200

    async def __aenter__(self):
        return self

    async def __aexit__(self, *_exc):
        return False

    async def json(self, content_type=None):
        return self.payload

    async def read(self):
        return b"{}"


class DeckChunkingTests(unittest.TestCase):
    def test_oversized_section_is_not_appended_to_the_previous_chunk(self):
        first = "# بخش الف\nیک سطر کوتاه"
        second = "# بخش ب\n" + "\n".join(
            f"line {index} " + "x" * 80 for index in range(60)
        )
        chunks = _chunk_long_text(first + "\n" + second, max_chars=500)

        self.assertTrue(all(len(chunk) <= 500 for chunk in chunks), chunks)
        # The short section stays byte-identical: its text must never absorb
        # lines from the following (oversized) section.
        self.assertEqual(chunks[0], first)


class ModelValidationTests(unittest.IsolatedAsyncioTestCase):
    """A model may *arrange* the user's text; it may never rewrite it."""

    def settings(self):
        return make_settings()

    async def test_a_valid_plan_is_accepted(self):
        session = _FakeLLMSession(
            {
                "sections": [
                    {"title": "بخش اول", "lines": ["این متن عیناً در سند می‌آید."]},
                ]
            }
        )
        slides = await plan_deck(
            SAMPLE, settings=self.settings(), session=session
        )
        assert slides is not None
        self.assertEqual(slides[0].title, "بخش اول")
        self.assertEqual(slides[0].bullets, ("این متن عیناً در سند می‌آید.",))
        self.assertEqual(slides[0].notes, "این متن عیناً در سند می‌آید.")

        document = build_office_document(SAMPLE)
        presentation = Presentation(
            io.BytesIO(build_office_pptx(document, slides))
        )
        self.assertIn(
            "این متن عیناً در سند می‌آید.",
            presentation.slides[2].notes_slide.notes_text_frame.text,
        )

    async def test_a_line_that_is_not_in_the_source_is_dropped(self):
        session = _FakeLLMSession(
            {
                "sections": [
                    {
                        "title": "بخش اول",
                        "lines": [
                            "این متن عیناً در سند می‌آید.",
                            "جمله‌ای که کاربر ننوشته است.",
                        ],
                    }
                ]
            }
        )
        slides = await plan_deck(
            SAMPLE, settings=self.settings(), session=session
        )
        assert slides is not None
        self.assertEqual(slides[0].bullets, ("این متن عیناً در سند می‌آید.",))

    async def test_a_plan_that_invents_everything_is_refused(self):
        session = _FakeLLMSession(
            {
                "sections": [
                    {"title": "بخش", "lines": ["یک", "دو", "سه"]},
                ]
            }
        )
        slides = await plan_deck(
            SAMPLE, settings=self.settings(), session=session
        )
        self.assertIsNone(slides)

    async def test_a_table_with_invented_cells_keeps_only_the_real_rows(self):
        session = _FakeLLMSession(
            {
                "headers": ["نام", "سن"],
                "rows": [
                    ["علی", "۲۰"],
                    ["کسی که نیست", "۹۹"],
                ],
            }
        )
        table = await infer_table(
            SAMPLE, settings=self.settings(), session=session
        )
        assert table is not None
        self.assertEqual(table.headers, ("نام", "سن"))
        self.assertEqual(table.rows, (("علی", "۲۰"),))

    async def test_a_table_of_invented_cells_is_refused(self):
        session = _FakeLLMSession(
            {"headers": ["نام"], "rows": [["کسی که نیست"]]}
        )
        table = await infer_table(
            SAMPLE, settings=self.settings(), session=session
        )
        self.assertIsNone(table)

    async def test_without_a_provider_the_planner_is_skipped(self):
        settings = make_settings(
            note_api_key="", note_api_base_url="", nara_api_key="", openrouter_api_key=""
        )
        session = _FakeLLMSession({})
        self.assertIsNone(await plan_deck(SAMPLE, settings=settings, session=session))
        self.assertIsNone(await infer_table(SAMPLE, settings=settings, session=session))




if __name__ == "__main__":
    unittest.main()
