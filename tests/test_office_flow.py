"""End-to-end tests for the Free Office menu, export and document-note flows."""

from __future__ import annotations

import hashlib
import io
import tempfile
import time as real_time
import unittest
import unicodedata
from dataclasses import replace
from pathlib import Path
from unittest import mock

import openpyxl
from pptx import Presentation
from pypdf import PdfReader

from gamas_bot.bot import GamasBot
from gamas_bot.db import Database
from gamas_bot.doc_extract import DocChunk, chunks_to_json
from gamas_bot.textnorm import normalize_for_compare
from tests.support import FakeJobEvent, FakeTelegramMessage, docx_text, make_settings


class CallbackEvent:
    """Private callback double with message ids and screen recording."""

    is_private = True

    def __init__(self, client, data: bytes, user_id: int, message_id: int = 77):
        self.client = client
        self.data = data
        self.sender_id = int(user_id)
        self.chat_id = int(user_id)
        self.message_id = int(message_id)
        self.message = FakeTelegramMessage(
            client, chat_id=self.chat_id, message_id=self.message_id
        )
        self.edits: list[dict] = []
        self.responses: list[dict] = []
        self.answers: list[tuple[str | None, bool]] = []

    async def answer(self, text=None, *, alert=False):
        self.answers.append((text, alert))

    async def edit(self, text, buttons=None):
        self.edits.append({"text": text, "buttons": buttons})

    async def respond(self, text="", buttons=None, **kwargs):
        self.responses.append({"text": text, "buttons": buttons, **kwargs})


class PlainMessageEvent:
    """A small private incoming-text event for the ordinary message router."""

    is_private = True

    def __init__(self, client, user_id: int, text: str):
        self.client = client
        self.sender_id = int(user_id)
        self.sender = None
        self.message = FakeTelegramMessage(
            client,
            chat_id=self.sender_id,
            message_id=88,
            text=text,
        )

    async def reply(self, text="", **kwargs):
        return await self.message.reply(text, **kwargs)

    async def respond(self, text="", **kwargs):
        return await self.message.respond(text, **kwargs)


class OfficeFlowTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.user_id = 5_501
        self.settings = make_settings(
            db_path=self.root / "bot.sqlite3",
            temp_dir=self.root / "private-tmp",
            office_enabled=True,
            office_min_interval_seconds=0.0,
            note_api_key="",
            note_api_base_url="",
            note_providers=(),
        )
        self.settings.temp_dir.mkdir(parents=True, exist_ok=True)
        self.db = Database(self.settings.db_path)
        await self.db.connect()
        await self.db.ensure_user(self.user_id)
        self.capture = FakeJobEvent(chat_id=self.user_id)
        self.bot = GamasBot(self.settings)
        self.bot.db = self.db
        self.bot.client = self.capture.client

    async def asyncTearDown(self):
        await self.bot.shutdown()
        await self.db.close()
        self.temp.cleanup()

    def callback(self, data: str, *, user_id: int | None = None, client=None) -> CallbackEvent:
        return CallbackEvent(
            client or self.capture.client,
            data.encode("ascii"),
            self.user_id if user_id is None else user_id,
        )

    async def billing_snapshot(self):
        cursor = await self.db.read_db.execute(
            "SELECT kind, seconds, note FROM ledger WHERE user_id = ? ORDER BY id",
            (self.user_id,),
        )
        ledger = tuple(tuple(row) for row in await cursor.fetchall())
        return (
            await self.db.available_seconds(self.user_id),
            await self.db.free_use_state(self.user_id),
            ledger,
        )

    async def enter_office(self) -> None:
        await self.bot._handle_callback(self.callback("office:start"))
        self.assertIn(self.user_id, self.bot._office_sessions)

    async def test_pasted_messages_make_real_word_excel_powerpoint_and_selectable_pdf_without_billing(self):
        before = await self.billing_snapshot()
        await self.enter_office()
        source = (
            "# گزارش آزمون\n"
            "این یک متن آزمایشی برای بررسی خروجی‌های واقعی Office است.\n"
            "این جمله باید در هر قالب باقی بماند.\n"
            "Office selectable proof token 781.\n"
            "نام\tمقدار\n"
            "علی\t۲۰\n"
        )
        incoming = PlainMessageEvent(self.capture.client, self.user_id, source)
        await self.bot._handle_message(incoming)
        self.assertIn(source.strip(), self.bot._office_sessions[self.user_id].text)

        for kind, suffix in (
            ("word", ".docx"),
            ("excel", ".xlsx"),
            ("ppt", ".pptx"),
            ("pdf", ".pdf"),
        ):
            callback = self.callback(f"office:fmt:{kind}")
            await self.bot._handle_callback(callback)
            self.assertTrue(any(path.endswith(suffix) for path, _payload in self.capture.file_payloads))
            self.assertTrue(callback.edits)
            self.assertIn("فایل ساخته شد", callback.edits[-1]["text"])

        word_bytes = self.capture.file_bytes(".docx")
        word_text = docx_text(word_bytes)
        self.assertIn("متن آزمایشی", word_text)
        self.assertIn("علی", word_text)

        workbook = openpyxl.load_workbook(
            io.BytesIO(self.capture.file_bytes(".xlsx")), data_only=False
        )
        all_cells = [str(value) for sheet in workbook.worksheets for row in sheet.values for value in row if value]
        self.assertTrue(any("متن آزمایشی" in value for value in all_cells))
        self.assertIn("علی", all_cells)

        presentation = Presentation(io.BytesIO(self.capture.file_bytes(".pptx")))
        slide_text = "\n".join(
            shape.text
            for slide in presentation.slides
            for shape in slide.shapes
            if getattr(shape, "has_text_frame", False)
        )
        self.assertIn("متن آزمایشی", slide_text)
        self.assertIn("علی", slide_text)

        pdf_reader = PdfReader(io.BytesIO(self.capture.file_bytes(".pdf")))
        pdf_text = "\n".join(page.extract_text() or "" for page in pdf_reader.pages)
        normalized_pdf = normalize_for_compare(unicodedata.normalize("NFKC", pdf_text))
        self.assertIn("office selectable proof token 781", normalized_pdf)
        self.assertIn("علی", normalized_pdf)

        after = await self.billing_snapshot()
        self.assertEqual(after, before)
        items = await self.db.library_list(self.user_id, limit=10)
        self.assertEqual(len(items), 4)
        self.assertTrue(all(item["kind"] == "office" for item in items))
        self.assertEqual(list(self.settings.temp_dir.iterdir()), [])

    async def test_office_build_shows_a_bar_and_saves_a_short_library_title(self):
        """The build is a bar from the first frame, and the row is a short name."""
        from gamas_bot.titles import word_count

        await self.enter_office()
        source = (
            "گزارش فروش ماهانه فروردین\n"
            "این متن برای بررسی نوار پیشرفت و عنوان کتابخانه است."
        )
        await self.bot._handle_message(
            PlainMessageEvent(self.capture.client, self.user_id, source)
        )
        callback = self.callback("office:fmt:word")
        await self.bot._handle_callback(callback)

        frames = [edit["text"] for edit in callback.edits]
        self.assertTrue(frames)
        # A bar exists before the file is built (the 4 % frame) and the final
        # screen still carries the outcome on a full bar.
        self.assertIn("٪", frames[0])
        self.assertTrue(any("٪" in frame for frame in frames))
        self.assertIn("فایل ساخته شد", frames[-1])
        self.assertIn("۱۰۰٪", frames[-1])

        items = await self.db.library_list(self.user_id, limit=5)
        self.assertEqual(len(items), 1)
        self.assertEqual(items[0]["kind"], "office")
        self.assertGreaterEqual(word_count(items[0]["title"]), 2)
        self.assertLessEqual(word_count(items[0]["title"]), 4)
        self.assertIn("Word", items[0]["title"])
        # The file itself is named by the same title as its library row: the
        # row, the cover and the file name cannot disagree.
        base = items[0]["title"].split(" — ")[0]
        delivered = [str(record["path"]) for record in self.capture.deliveries]
        self.assertTrue(any(base in path for path in delivered), delivered)

    async def test_uploaded_csv_is_parsed_and_formula_like_cells_are_literal(self):
        source_file = self.root / "report.csv"
        source_file.write_text(
            'name,value\n"=1+2","comma, value"\n', encoding="utf-8"
        )
        upload = FakeJobEvent(source=source_file, chat_id=self.user_id)
        upload.message.attach_document(
            "report.csv", mime_type="text/csv", size=source_file.stat().st_size
        )
        self.bot.client = upload.client
        before = await self.billing_snapshot()

        await self.bot._handle_message(upload)
        session = self.bot._office_sessions[self.user_id]
        self.assertEqual(session.filename, "report.csv")
        self.assertIn('"=1+2","comma, value"', session.text)

        callback = self.callback("office:fmt:excel", client=upload.client)
        await self.bot._handle_callback(callback)
        workbook = openpyxl.load_workbook(
            io.BytesIO(upload.file_bytes(".xlsx")), data_only=False
        )
        sheet = workbook.active
        self.assertEqual(sheet["A2"].value, "=1+2")
        self.assertEqual(sheet["A2"].data_type, "s")
        self.assertIsNone(sheet["A2"].hyperlink)
        self.assertEqual(sheet["B2"].value, "comma, value")
        self.assertTrue(source_file.exists())
        self.assertEqual(
            source_file.read_text(encoding="utf-8"),
            'name,value\n"=1+2","comma, value"\n',
        )
        self.assertEqual(await self.billing_snapshot(), before)
        self.assertEqual(list(self.settings.temp_dir.iterdir()), [])

    async def test_a_text_upload_with_control_characters_still_exports_excel(self):
        """A verbose ``.txt`` may carry VT/FF; the Excel export must still build.

        ``read_text_file`` accepts every non-NUL byte of a text upload, so a
        form feed from a printer dump or a stray ``\\x01`` reached the builders.
        Excel refused the whole file (``IllegalCharacterError``) and the generic
        callback error was the only thing the user ever saw, while Word,
        PowerPoint and PDF of the same text worked.
        """
        source_text = (
            "گزارش کنترل\x01نویسه\n"
            "این سطر \x0c باید با فاصله جدا بماند\n"
            "نام\tمقدار\n"
            "علی\t۲۰\n"
        )
        source_file = self.root / "verbose.txt"
        source_file.write_text(source_text, encoding="utf-8")
        upload = FakeJobEvent(source=source_file, chat_id=self.user_id)
        upload.message.attach_document(
            "verbose.txt", mime_type="text/plain", size=source_file.stat().st_size
        )
        self.bot.client = upload.client
        before = await self.billing_snapshot()

        await self.bot._handle_message(upload)
        self.assertIn(self.user_id, self.bot._office_sessions)

        callback = self.callback("office:fmt:excel", client=upload.client)
        await self.bot._handle_callback(callback)
        self.assertTrue(
            any(path.endswith(".xlsx") for path, _payload in upload.file_payloads),
            f"no Excel file was delivered: {upload.file_paths}",
        )
        self.assertIn("فایل ساخته شد", callback.edits[-1]["text"])

        workbook = openpyxl.load_workbook(
            io.BytesIO(upload.file_bytes(".xlsx")), data_only=False
        )
        cells = [
            str(value)
            for sheet in workbook.worksheets
            for row in sheet.values
            for value in row
            if value
        ]
        joined = "\n".join(cells)
        self.assertIn("گزارش کنترل", joined)
        self.assertIn("نویسه", joined)
        # The soft line break still separates the two words.
        self.assertIn("این سطر", joined)
        self.assertIn("باید با فاصله جدا بماند", joined)
        self.assertIn("علی", cells)
        for value in cells:
            self.assertFalse(
                any(ord(char) < 0x20 and char not in "\t\n\r" for char in value),
                f"a control character reached the sheet: {value!r}",
            )
        # The source file is the user's; it is never rewritten.
        self.assertEqual(source_file.read_text(encoding="utf-8"), source_text)
        self.assertEqual(await self.billing_snapshot(), before)
        self.assertEqual(list(self.settings.temp_dir.iterdir()), [])

    async def test_first_export_is_not_throttled_when_the_host_just_booted(self):
        """Audit F-3: ``0.0`` is a real ``time.monotonic()`` value, not "never".

        The per-user throttle read ``self._office_last_use.get(user_id, 0.0)``,
        so a host whose uptime was below ``OFFICE_MIN_INTERVAL_SECONDS`` refused
        the very first export of every user and told them to wait. The fake
        clock is injected into ``gamas_bot.bot`` only, so the event loop keeps
        its own real clock.
        """

        class FreshBootClock:
            value = 5.0

            def monotonic(self):
                return self.value

            def __getattr__(self, name):
                return getattr(real_time, name)

        self.settings = replace(self.settings, office_min_interval_seconds=10.0)
        self.bot.settings = self.settings
        self.bot._office_last_use.clear()
        fresh_clock = FreshBootClock()

        await self.enter_office()
        self.bot._office_add_source(
            self.bot._office_sessions[self.user_id], "متن آزمایشی برای شمارش زمان"
        )
        with mock.patch("gamas_bot.bot.time", fresh_clock):
            first = self.callback("office:fmt:word")
            await self.bot._handle_callback(first)
        texts = [edit["text"] for edit in first.edits]
        throttle_seen = [
            text
            for text, _alert in first.answers
            if "ثانیه دیگر دوباره تلاش کنید" in (text or "")
        ]
        self.assertEqual(
            throttle_seen, [], "a fresh boot must not throttle the first export"
        )
        self.assertTrue(any("فایل ساخته شد" in text for text in texts))

        # The throttle itself is unchanged: an immediate second export waits.
        with mock.patch("gamas_bot.bot.time", fresh_clock):
            second = self.callback("office:fmt:word")
            await self.bot._handle_callback(second)
        self.assertTrue(
            any(
                "ثانیه دیگر دوباره تلاش کنید" in (text or "")
                for text, _alert in second.answers
            ),
            "the per-user throttle must still apply after the first export",
        )

    async def test_uploaded_document_gets_full_and_compact_word_pdf_with_source_locations(self):
        source = (
            "سامانه از چند بخش مستقل تشکیل شده است.\n"
            "هر بخش مسئولیت مشخصی دارد و پیام‌ها به ترتیب پردازش می‌شوند.\n"
            "تنظیمات در یک فایل مرکزی نگهداری می‌شوند.\n"
            "کاربر می‌تواند خروجی Word یا PDF دریافت کند.\n"
            "فایل اصلی هنگام ساخت خروجی تغییر نمی‌کند."
        )
        locator = "صفحهٔ ۱"
        cached = [DocChunk(0, locator, (locator,), f"[{locator}]\n{source}")]
        doc_id, created = await self.db.document_add(
            self.user_id,
            file_msg_ref=f"{self.user_id}:900",
            title="راهنمای سامانه.pdf",
            source_type="pdf",
            page_count=1,
            sha256=hashlib.sha256(source.encode("utf-8")).hexdigest(),
            text=source,
            chunks_json=chunks_to_json(cached),
        )
        self.assertTrue(created)
        before = await self.billing_snapshot()

        full = self.callback(f"doc:note:{doc_id}:full")
        await self.bot._handle_callback(full)
        compact = self.callback(f"doc:note:{doc_id}:compact")
        await self.bot._handle_callback(compact)

        delivered = self.capture.file_payloads
        full_word = next(payload for path, payload in delivered if "FULL" in path and path.endswith(".docx"))
        compact_word = next(payload for path, payload in delivered if "COMPACT" in path and path.endswith(".docx"))
        full_pdf = next(payload for path, payload in delivered if "FULL" in path and path.endswith(".pdf"))
        compact_pdf = next(payload for path, payload in delivered if "COMPACT" in path and path.endswith(".pdf"))
        full_text = docx_text(full_word)
        compact_text = docx_text(compact_word)
        self.assertIn("منبع: صفحهٔ ۱", full_text)
        self.assertIn("منبع: صفحهٔ ۱", compact_text)
        self.assertIn("پیام‌ها به ترتیب پردازش می‌شوند", full_text)
        self.assertNotEqual(full_text, compact_text)
        self.assertLess(len(compact_text), len(full_text))
        for pdf_bytes in (full_pdf, compact_pdf):
            self.assertTrue(pdf_bytes.startswith(b"%PDF-"))
            extracted = "\n".join(
                page.extract_text() or "" for page in PdfReader(io.BytesIO(pdf_bytes)).pages
            )
            self.assertIn(
                "صفحه",
                normalize_for_compare(unicodedata.normalize("NFKC", extracted)),
            )

        stored = await self.db.document_get(self.user_id, doc_id)
        self.assertEqual(stored["text"], source)
        self.assertEqual(stored["file_msg_ref"], f"{self.user_id}:900")
        self.assertEqual(await self.billing_snapshot(), before)
        items = await self.db.library_list(self.user_id, limit=10)
        self.assertEqual(len(items), 4)
        self.assertTrue(all(item["kind"] == "office" for item in items))
        # A document-note booklet is named by the same short title: the cover,
        # both files and all four rows share it.
        from gamas_bot.titles import word_count

        for item in items:
            self.assertGreaterEqual(word_count(item["title"]), 2, item["title"])
            self.assertLessEqual(word_count(item["title"]), 4, item["title"])
        base = items[0]["title"].split(" — ")[0]
        self.assertEqual(base, "راهنمای سامانه")
        self.assertIn(base, full_text)
        self.assertIn(base, compact_text)
        self.assertNotIn("راهنمای سامانه.pdf", items[0]["title"])

    async def test_foreign_document_callback_is_owner_scoped(self):
        other = self.user_id + 1
        await self.db.ensure_user(other)
        source = "فقط مالک باید این سند را دریافت کند."
        doc_id, _created = await self.db.document_add(
            other,
            file_msg_ref=f"{other}:900",
            title="خصوصی.pdf",
            source_type="pdf",
            page_count=1,
            sha256=hashlib.sha256(source.encode("utf-8")).hexdigest(),
            text=source,
            chunks_json="[]",
        )
        callback = self.callback(f"doc:ex:{doc_id}:word")
        await self.bot._handle_callback(callback)
        self.assertIn(("این سند پیدا نشد.", True), callback.answers)
        self.assertEqual(self.capture.file_payloads, [])

    async def test_disabled_office_hides_menu_and_rejects_stale_callback(self):
        self.bot.settings = replace(self.settings, office_enabled=False)
        payloads = {
            button.type.data
            for row in self.bot._user_menu()
            for button in row
        }
        self.assertNotIn(b"office:start", payloads)
        callback = self.callback("office:start")
        await self.bot._handle_callback(callback)
        self.assertEqual(callback.answers[-1], ("آفیس رایگان در تنظیمات سرور غیرفعال است.", True))
        self.assertNotIn(self.user_id, self.bot._office_sessions)


if __name__ == "__main__":
    unittest.main()
