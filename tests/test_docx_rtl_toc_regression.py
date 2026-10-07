"""Regressions for the real-render RTL and static-TOC fixes.

Covers the bugs found by rendering delivered documents with LibreOffice:

* bidi paragraphs carried *physical* jc/ind values; OOXML defines them as
  logical (start/end) in RTL paragraphs, so headings rendered on the left;
* Latin-only paragraphs wrote ``<w:bidi/>`` when they meant LTR;
* TOC page numbers were mapped from link annotations that renderers may not
  emit; they now come from the heading outline and fail closed;
* the dot leader of a trailing tab was not painted;
* bookmark names were not per-document deterministic.
"""

from __future__ import annotations

import io
import re
import unittest
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path
from unittest.mock import patch

from docx import Document

from gamas_bot import docx_export
from gamas_bot.docx_export import (
    BOOKMARK_NAME_RE,
    BOOKMARK_PREFIX,
    TOC_LEADER_TERMINATOR,
    TOC_TITLE,
    DocumentMeta,
    DocxPaginationError,
    _enable_bidi,
    _rendered_toc_page_numbers,
    _title_key,
    build_notes_docx,
    normalize_bidi_alignment,
    resolve_design,
)
from gamas_bot.structuring import NoteSection, StructuredNotes

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
META = DocumentMeta(reference="GMS-REG")


def _paragraph_docx(p_pr: str, text: str = "متن") -> bytes:
    """A minimal DOCX whose single paragraph has the given pPr XML."""
    document = Document()
    document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    source = zipfile.ZipFile(io.BytesIO(buffer.getvalue()))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as target:
        for item in source.infolist():
            data = source.read(item)
            if item.filename == "word/document.xml":
                data = re.sub(rb"<w:p>|<w:p [^>]*>", b"<w:p>" + p_pr.encode(), data, count=1)
            target.writestr(item, data)
    return out.getvalue()


def _first_ppr(data: bytes) -> ET.Element:
    xml = zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml")
    return ET.fromstring(xml).find(f".//{W}p/{W}pPr")


def _notes(titles: list[str]) -> StructuredNotes:
    body = "این یک پاراگراف آزمایشی دربارهٔ HbA1c و 500 mg است. " * 30
    return StructuredNotes(
        title="جزوهٔ آزمایشی",
        summary="خلاصه",
        sections=tuple(NoteSection(heading=t, paragraphs=(body,)) for t in titles),
    )


class BidiAlignmentTests(unittest.TestCase):
    def test_rtl_paragraph_alignment_and_indents_are_logical(self):
        data = normalize_bidi_alignment(_paragraph_docx(
            '<w:pPr><w:bidi/><w:ind w:right="567" w:hanging="200"/><w:jc w:val="right"/></w:pPr>'
        ))
        p_pr = _first_ppr(data)
        self.assertEqual(p_pr.find(f"{W}jc").get(f"{W}val"), "left")
        ind = p_pr.find(f"{W}ind")
        self.assertEqual(ind.get(f"{W}left"), "567")
        self.assertIsNone(ind.get(f"{W}right"))
        self.assertEqual(ind.get(f"{W}hanging"), "200")

    def test_ltr_paragraphs_are_left_untouched(self):
        for p_pr in (
            '<w:pPr><w:ind w:right="567"/><w:jc w:val="right"/></w:pPr>',
            '<w:pPr><w:bidi w:val="0"/><w:ind w:right="567"/><w:jc w:val="right"/></w:pPr>',
        ):
            with self.subTest(p_pr=p_pr):
                parsed = _first_ppr(normalize_bidi_alignment(_paragraph_docx(p_pr)))
                self.assertEqual(parsed.find(f"{W}jc").get(f"{W}val"), "right")
                self.assertEqual(parsed.find(f"{W}ind").get(f"{W}right"), "567")

    def test_center_and_both_are_direction_neutral(self):
        for value in ("center", "both"):
            with self.subTest(value=value):
                data = normalize_bidi_alignment(
                    _paragraph_docx(f'<w:pPr><w:bidi/><w:jc w:val="{value}"/></w:pPr>')
                )
                self.assertEqual(_first_ppr(data).find(f"{W}jc").get(f"{W}val"), value)

    def test_ltr_paragraph_writes_explicit_bidi_off(self):
        document = Document()
        paragraph = document.add_paragraph("HbA1c = 6.5%")
        _enable_bidi(paragraph, rtl=False)
        bidi = paragraph._p.pPr.find(f"{W}bidi")
        self.assertIsNotNone(bidi)
        self.assertEqual(bidi.get(f"{W}val"), "0")

    def test_generated_rtl_headings_are_visually_right_aligned(self):
        data = build_notes_docx(_notes(["الف", "ب"]), meta=META,
                                design=resolve_design({"toc_enabled": False}))
        root = ET.fromstring(zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml"))
        checked = 0
        for p_pr in root.iter(f"{W}pPr"):
            style = p_pr.find(f"{W}pStyle")
            bidi = p_pr.find(f"{W}bidi")
            if style is None or not style.get(f"{W}val", "").startswith("Heading"):
                continue
            if bidi is None or bidi.get(f"{W}val") in ("0", "false"):
                continue
            jc = p_pr.find(f"{W}jc")
            if jc is not None:
                checked += 1
                self.assertNotEqual(jc.get(f"{W}val"), "right")
        self.assertGreater(checked, 0)


class BookmarkTests(unittest.TestCase):
    def _bookmarks(self, data: bytes) -> list[str]:
        root = ET.fromstring(zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml"))
        return [b.get(f"{W}name") for b in root.iter(f"{W}bookmarkStart")]

    def test_names_are_sanitized_unique_and_deterministic_with_duplicate_titles(self):
        notes = _notes(["مقدمه", "روش‌ها", "مقدمه", "روش‌ها", "مقدمه ' \" <x>"])
        design = resolve_design({"toc_enabled": False})
        first = self._bookmarks(build_notes_docx(notes, meta=META, design=design))
        second = self._bookmarks(build_notes_docx(notes, meta=META, design=design))
        self.assertGreaterEqual(len(first), 5)
        self.assertEqual(len(first), len(set(first)))
        self.assertEqual(first, second)  # per-document sequence, not a global counter
        for name in first:
            self.assertRegex(name, BOOKMARK_NAME_RE)
            self.assertTrue(name.startswith(BOOKMARK_PREFIX))
            self.assertTrue(name.startswith("_"))  # hidden from Word's bookmark list


class FakePage:
    def __init__(self, text: str):
        self.text = text

    def extract_text(self):
        return self.text

    def get(self, key, default=None):
        return default


class FakeReader:
    """Stands in for pypdf.PdfReader: pages plus a heading outline."""

    outline_items: list = []
    page_texts: list = []

    def __init__(self, *_args, **_kwargs):
        self.pages = [FakePage(text) for text in self.page_texts]
        self.outline = [{"/Title": title, "_page": page} for title, page in self.outline_items]

    def get_destination_page_number(self, item):
        return item["_page"] - 1


class OutlinePageMappingTests(unittest.TestCase):
    ENTRIES = [
        {"anchor": "_GamasH1", "title": "۱. مقدمه", "level": 1},
        {"anchor": "_GamasH2", "title": "۲. روش‌ها", "level": 1},
        {"anchor": "_GamasH3", "title": "۳. مقدمه", "level": 1},
    ]

    def measure(self, outline, pages=("جلد", TOC_TITLE, "", "", "", "")):
        FakeReader.outline_items = outline
        FakeReader.page_texts = list(pages)
        with patch("pypdf.PdfReader", FakeReader), patch.object(
            docx_export, "_render_pdf", return_value=Path("/nonexistent.pdf")
        ):
            return _rendered_toc_page_numbers(
                b"docx", self.ENTRIES, timeout_seconds=5, renderer_bin="/bin/true"
            )

    def test_pages_come_from_the_outline_in_order(self):
        outline = [("جزوهٔ آزمایشی", 3), ("۱. مقدمه", 3), ("۲. روش‌ها", 4), ("۳. مقدمه", 6)]
        self.assertEqual(self.measure(outline), [3, 4, 6])

    def test_titles_match_across_digit_and_letter_variants(self):
        outline = [("1. مقدمه", 3), ("2. روش\u200cها", 5), ("3.  مقدمه", 5)]
        self.assertEqual(self.measure(outline), [3, 5, 5])
        self.assertEqual(_title_key("كتاب ي ۱"), _title_key("کتاب ی 1"))

    def test_reversed_toc_title_text_is_accepted(self):
        outline = [("۱. مقدمه", 3), ("۲. روش‌ها", 4), ("۳. مقدمه", 5)]
        pages = ("جلد", TOC_TITLE[::-1], "", "", "", "")
        self.assertEqual(self.measure(outline, pages), [3, 4, 5])

    def test_fail_closed_cases(self):
        good = [("۱. مقدمه", 3), ("۲. روش‌ها", 4), ("۳. مقدمه", 5)]
        cases = {
            "missing heading": ([("۱. مقدمه", 3), ("۲. روش‌ها", 4)], None),
            "heading before page 3": ([("۱. مقدمه", 2), ("۲. روش‌ها", 4), ("۳. مقدمه", 5)], None),
            "no TOC on page 2": (good, ("جلد", "متن", "", "", "")),
            "TOC on the cover": (good, (TOC_TITLE, TOC_TITLE, "", "", "")),
            "too few pages": (good, ("جلد", TOC_TITLE)),
            "beyond the last page": ([("۱. مقدمه", 3), ("۲. روش‌ها", 4), ("۳. مقدمه", 9)], None),
        }
        for name, (outline, pages) in cases.items():
            with self.subTest(name):
                with self.assertRaises(DocxPaginationError):
                    if pages is None:
                        self.measure(outline)
                    else:
                        self.measure(outline, pages)

    def test_missing_renderer_fails_explicitly(self):
        with patch.object(docx_export.shutil, "which", return_value=None):
            with self.assertRaises(DocxPaginationError):
                _rendered_toc_page_numbers(b"docx", self.ENTRIES, timeout_seconds=5)


class TocMarkupTests(unittest.TestCase):
    def _build(self, titles):
        with patch.object(
            docx_export, "_rendered_toc_page_numbers",
            side_effect=lambda data, entries, **_: [3 + i for i in range(len(entries))],
        ):
            return build_notes_docx(_notes(titles), meta=META,
                                    design=resolve_design({"toc_enabled": True}))

    def test_rows_have_a_painted_dot_leader_and_clickable_numbers(self):
        data = self._build([f"موضوع {i}" for i in range(1, 11)])
        parts = zipfile.ZipFile(io.BytesIO(data))
        root = ET.fromstring(parts.read("word/document.xml"))
        self.assertNotIn("updateFields", parts.read("word/settings.xml").decode())
        table = next(child for child in root.find(f"{W}body") if child.tag == f"{W}tbl")
        rows = table.findall(f"{W}tr")
        self.assertGreaterEqual(len(rows), 10)
        persian = str.maketrans("0123456789", "۰۱۲۳۴۵۶۷۸۹")
        for index, row in enumerate(rows):
            title_cell, page_cell = row.findall(f"{W}tc")
            tab = title_cell.find(f".//{W}tabs/{W}tab[@{W}leader='dot']")
            self.assertIsNotNone(tab)
            # Logical end stop of a bidi paragraph (visual left).
            self.assertEqual(tab.get(f"{W}val"), "right")
            runs = title_cell.find(f"{W}p").findall(f"{W}r")
            tab_run = runs[-1]
            self.assertIsNotNone(tab_run.find(f"{W}tab"))
            # A leader is only painted when something follows the tab.
            self.assertEqual(tab_run.find(f"{W}t").text, TOC_LEADER_TERMINATOR)
            # Level with the last line of a wrapped title (where the leader ends).
            self.assertEqual(page_cell.find(f"{W}tcPr/{W}vAlign").get(f"{W}val"), "bottom")
            page_link = page_cell.find(f".//{W}hyperlink")
            self.assertIsNotNone(page_link)
            self.assertEqual(
                "".join(t.text or "" for t in page_link.iter(f"{W}t")),
                str(3 + index).translate(persian),
            )
            self.assertIsNone(page_link.find(f".//{W}u[@{W}val='single']"))

    def test_toc_has_no_field_or_placeholder_left(self):
        data = self._build([f"موضوع {i}" for i in range(1, 6)])
        xml = zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml").decode()
        self.assertNotIn("instrText", xml)
        self.assertNotIn("fldChar", xml)
        root = ET.fromstring(xml)
        table = next(child for child in root.find(f"{W}body") if child.tag == f"{W}tbl")
        for row in table.findall(f"{W}tr"):
            page_text = "".join(t.text or "" for t in row.findall(f"{W}tc")[1].iter(f"{W}t"))
            self.assertNotEqual(page_text, docx_export.TOC_PAGE_PLACEHOLDER)
            self.assertRegex(page_text, r"^[۰-۹]+$")


if __name__ == "__main__":
    unittest.main()
