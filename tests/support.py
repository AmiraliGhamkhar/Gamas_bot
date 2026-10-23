"""Shared helpers for the test-suite: settings factory, deck builder, media.

Media payloads are *real* files (generated with the standard library and PyAV)
because the bot now processes media through the PyAV-based worker instead of
stub binaries — every fixture the suite probes/merges must actually decode.
"""

from __future__ import annotations

import inspect
import io
import tempfile
import wave
import zipfile
from dataclasses import replace
from pathlib import Path

from telethon import TelegramClient

from gamas_bot.config import (
    DEFAULT_CEREBRAS_BASE_URL,
    DEFAULT_CEREBRAS_MODEL,
    DEFAULT_GEMINI_MODEL,
    DEFAULT_GEMINI_THINKING_BUDGET,
    DEFAULT_GROQ_BASE_URL,
    DEFAULT_GROQ_MODEL,
    Settings,
)

AUDIO_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/audio"
VIDEO_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/video"
RELS_NS = "http://schemas.openxmlformats.org/package/2006/relationships"

# The committed legacy-deck fixtures come from the MIT-licensed ppt2pptx test
# corpus (https://github.com/HuiTurn/ppt2pptx); see tests/fixtures/README.md.
FIXTURES_DIR = Path(__file__).resolve().parent / "fixtures"


class FakeTelegramClient:
    """A ``TelegramClient`` double that enforces the *installed* Telethon API.

    A permissive double is what hid the production failure: ``Message.reply``
    forwards to ``send_message``, which has **no** ``caption`` parameter (the
    caption of a file message is its ``message`` argument), so
    ``reply(file=…, caption=…)`` raises ``TypeError`` on a real client while a
    hand-written fake accepted it. Every call here is therefore bound against
    the real signature, and ``send_message`` forwards a file to ``send_file``
    exactly as Telethon does — an unsupported keyword fails the same way it
    fails in production.
    """

    SEND_MESSAGE = inspect.signature(TelegramClient.send_message)
    SEND_FILE = inspect.signature(TelegramClient.send_file)

    def __init__(self, recorder: "FakeJobEvent | None" = None, *, fail_file_delivery: bool = False):
        self.recorder = recorder
        self.fail_file_delivery = fail_file_delivery
        #: every text message, as the arguments it was sent with.
        self.texts: list[dict] = []
        #: every document/photo delivery, as the arguments it was sent with.
        self.deliveries: list[dict] = []
        self.disconnected = False

    async def disconnect(self) -> None:
        """The shutdown path calls this on the real client."""
        self.disconnected = True

    @staticmethod
    def _bind(signature, name: str, *args, **kwargs):
        """Bind a call to the real Telethon signature (``None`` is ``self``)."""
        try:
            return signature.bind(None, *args, **kwargs)
        except TypeError as exc:  # the error production would raise, verbatim
            raise TypeError(f"{name}() {exc}") from None

    async def send_message(self, *args, **kwargs):
        bound = self._bind(self.SEND_MESSAGE, "send_message", *args, **kwargs)
        bound.apply_defaults()
        arguments = dict(bound.arguments)
        entity = arguments.pop("entity")
        message = arguments.pop("message")
        file = arguments.pop("file", None)
        if file is not None:
            # Telethon: ``send_message`` delegates to ``send_file`` with the
            # message text as the caption (``MessageMethods.send_message``).
            forwarded = {
                key: value
                for key, value in arguments.items()
                if key in self.SEND_FILE.parameters and key != "reply_to"
            }
            return await self.send_file(
                entity, file, caption=message, reply_to=arguments.get("reply_to"), **forwarded
            )
        record = {"entity": entity, "text": message, **arguments}
        self.texts.append(record)
        if self.recorder is not None:
            self.recorder.replies.append(message)
        return FakeSentMessage(self, record)

    async def send_file(self, *args, **kwargs):
        bound = self._bind(self.SEND_FILE, "send_file", *args, **kwargs)
        bound.apply_defaults()
        arguments = dict(bound.arguments)
        entity = arguments.pop("entity")
        path = str(arguments.pop("file"))
        if self.fail_file_delivery:
            raise RuntimeError("Telegram refused the upload (simulated)")
        record = {"entity": entity, "path": path, **arguments}
        self.deliveries.append(record)
        if self.recorder is not None:
            caption = arguments.get("caption") or ""
            try:
                payload = Path(path).read_bytes()
            except OSError:
                payload = b""
            self.recorder.files.append((caption, path))
            self.recorder.file_payloads.append((path, payload))
        return FakeSentMessage(self, record)


class FakeSentMessage:
    """The message a send produced: enough of Telethon's ``Message`` to edit."""

    def __init__(self, client: FakeTelegramClient, record: dict):
        self.client = client
        self.record = record

    async def edit(self, text: str = "", buttons=None) -> "FakeSentMessage":
        self.record["text"] = text
        if self.client.recorder is not None:
            self.client.recorder.edits.append(text)
        return self


class FakeTelegramFile:
    """The ``message.file`` double: name, mime type, size and duration."""

    def __init__(self, name=None, mime_type=None, size=0, duration=None):
        self.name = name
        self.mime_type = mime_type
        self.size = size
        self.duration = duration


class FakeTelegramDocument:
    """The ``message.document`` double of an uploaded file."""

    def __init__(self, file: FakeTelegramFile):
        self.file = file
        self.attributes: list = []


class FakeTelegramMessage:
    """The incoming upload message, with Telethon's reply/respond contract.

    ``reply`` sets ``reply_to`` to this message and calls ``send_message`` on the
    client, exactly like ``telethon.tl.custom.message.Message.reply`` — so code
    that replies with an unsupported keyword fails here too.

    The media attributes (``voice``, ``audio``, ``document``, ``photo`` …) are
    the ones the bot classifies an upload by, so a message can be turned into
    an upload with :meth:`attach_document` or :meth:`attach_voice`.
    """

    def __init__(
        self,
        client: FakeTelegramClient,
        *,
        chat_id: int,
        message_id: int,
        download=None,
        text: str = "",
    ):
        self.client = client
        self._chat_id = chat_id
        self.id = message_id
        self._download = download
        #: The text of the message (Telethon calls it ``message``).
        self.message = text
        self.media = None
        self.file = FakeTelegramFile()
        self.document = None
        self.voice = None
        self.audio = None
        self.video = None
        self.video_note = None
        self.gif = None
        self.sticker = None
        self.photo = None

    def attach_document(self, name: str, *, mime_type: str | None = None, size: int = 0) -> None:
        """Make this message an uploaded file."""
        self.file = FakeTelegramFile(name, mime_type, size)
        self.document = FakeTelegramDocument(self.file)
        self.media = self.document

    def attach_voice(self, *, duration: float = 1.0) -> None:
        """Make this message a Telegram voice note."""
        self.file = FakeTelegramFile("voice.ogg", "audio/ogg", 0, duration)
        self.voice = self.file
        self.media = self.voice

    @property
    def chat_id(self) -> int:
        return self._chat_id

    async def get_input_chat(self):
        return self._chat_id

    async def reply(self, *args, **kwargs):
        kwargs["reply_to"] = self.id
        return await self.client.send_message(self._chat_id, *args, **kwargs)

    async def respond(self, *args, **kwargs):
        return await self.reply(*args, **kwargs)

    async def download_media(
        self, file: str | None = None, *, progress_callback=None
    ) -> str | None:
        if self._download is None:
            raise AssertionError("this fake message carries no downloadable media")
        return await self._download(file, progress_callback=progress_callback)


class FakeJobEvent:
    """Telethon event double wired to a signature-enforcing client.

    File payloads are snapshotted at send time because the job's temporary
    directory is removed right after delivery.
    """

    #: The bot only answers private chats.
    is_private = True

    def __init__(self, source: Path | None = None, *, chat_id: int = 555, message_id: int = 77):
        self.source = source
        self.sender_id = chat_id
        self.replies: list[str] = []
        self.responses: list[str] = []
        self.files: list[tuple[str, str]] = []
        self.file_payloads: list[tuple[str, bytes]] = []
        self.edits: list[str] = []
        self.client = FakeTelegramClient(self)
        self.message = FakeTelegramMessage(
            self.client,
            chat_id=chat_id,
            message_id=message_id,
            download=self._download if source is not None else None,
        )

    @property
    def deliveries(self) -> list[dict]:
        return self.client.deliveries

    async def _download(self, file: str, *, progress_callback=None) -> str:
        total = self.source.stat().st_size
        received = 0
        with self.source.open("rb") as source, open(file, "wb") as target:
            while chunk := source.read(64 * 1024):
                target.write(chunk)
                received += len(chunk)
                if progress_callback is not None:
                    # Match Telethon's (received_bytes, total_bytes) callback.
                    progress_callback(received, total if self.message.file.size else 0)
        return file

    async def reply(self, text: str = "", **kwargs) -> FakeSentMessage:
        return await self.message.reply(text, **kwargs)

    async def respond(self, text: str = "", **kwargs) -> FakeSentMessage:
        return await self.message.respond(text, **kwargs)

    @property
    def file_paths(self) -> list[str]:
        return [path for _caption, path in self.files]

    def file_bytes(self, suffix: str) -> bytes:
        """Content of the first sent file whose name ends with ``suffix``."""
        for path, payload in self.file_payloads:
            if path.endswith(suffix):
                return payload
        raise AssertionError(f"no sent file ends with {suffix!r}: {self.file_paths}")

    def delivery(self, suffix: str) -> dict:
        """The recorded ``send_file`` call whose path ends with ``suffix``."""
        for record in self.deliveries:
            if str(record["path"]).endswith(suffix):
                return record
        raise AssertionError(f"no send_file call delivered {suffix!r}: {self.file_paths}")


class FakeCallbackEvent:
    """A Telethon CallbackQuery double: the panels' buttons are answered with it.

    The panel is a screen tree, so the buttons have to be pressable in a
    test. ``data`` is the raw ``callback_data`` byte string, exactly as
    Telegram delivers it.
    """

    is_private = True

    def __init__(
        self,
        client: FakeTelegramClient,
        *,
        data: bytes,
        chat_id: int = 555,
        message_id: int = 77,
    ):
        self.client = client
        self.data = data
        self.sender_id = chat_id
        self.message = FakeTelegramMessage(client, chat_id=chat_id, message_id=message_id)
        self.answers: list[tuple[str | None, bool]] = []
        self.edits: list[dict] = []

    async def answer(self, text: str | None = None, alert: bool = False) -> None:
        self.answers.append((text, alert))

    async def respond(self, *args, **kwargs):
        return await self.message.respond(*args, **kwargs)

    async def reply(self, *args, **kwargs):
        return await self.message.reply(*args, **kwargs)

    async def edit(self, text: str = "", buttons=None):
        self.edits.append({"text": text, "buttons": buttons})
        return FakeSentMessage(self.client, {"text": text})


def sample_notes_json(title: str = "جزوهٔ آزمایشی") -> str:
    """A valid strict-JSON note answer, as the note API is now required to emit."""
    return (
        "{"
        f'"title": "{title}",'
        '"summary": "خلاصهٔ آزمایشی",'
        '"sections": [{'
        '"heading": "بخش نخست",'
        '"paragraphs": ["متن بخش نخست"],'
        '"bullets": ["نکتهٔ یک"],'
        '"key_points": ["نکتهٔ کلیدی"],'
        '"callouts": [{"kind": "نکته", "text": "یادآوری تست"}]'
        "}],"
        '"key_points": ["نکتهٔ کلیدی کل"],'
        '"glossary": [{"term": "HbA1c", "definition": "هموگلوبین گلیکوزیله"}]'
        "}"
    )


class FakeSessionContext:
    """Stand-in for aiohttp.ClientSession used as an async context manager."""

    async def __aenter__(self):
        return self

    async def __aexit__(self, *_exc):
        return False


def fake_session_factory(**_kwargs):
    return FakeSessionContext()


def docx_text(docx_bytes: bytes) -> str:
    """All paragraph and table text of a generated Word document."""
    import io

    import docx as docx_library

    document = docx_library.Document(io.BytesIO(docx_bytes))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def make_settings(**overrides) -> Settings:
    """A valid v3 Settings object; ``overrides`` may replace any field."""
    base = Settings(
        bot_token="token",
        api_id=1,
        api_hash="hash",
        admin_ids=frozenset({1}),
        session_path=Path("session"),
        db_path=Path("db.sqlite3"),
        temp_dir=Path("tmp"),
        log_file=None,
        log_level="INFO",
        log_format="plain",
        log_max_bytes=100_000,
        log_backup_count=1,
        max_file_size=2_000_000_000,
        max_active_jobs=2,
        max_active_jobs_per_user=8,
        max_queued_jobs=8,
        media_worker_concurrency=1,
        media_timeout=60.0,
        convert_timeout=60.0,
        stt_job_timeout=100.0,
        stt_poll_interval=1.0,
        temp_min_free_bytes=0,
        download_stall_seconds=300.0,
        audio_format_policy="auto",
        telegram_download_connections=1,
        presentation_max_clips=50,
        presentation_max_total_duration=3600.0,
        presentation_max_unpacked_bytes=1024**3,
        presentation_max_temp_bytes=8 * 1024**3,
        presentation_min_clip_seconds=0.5,
        presentation_silence_seconds=0.25,
        presentation_wav_limit_bytes=1024**3,
        presentation_include_slide_text=True,
        presentation_include_video_audio=True,
        presentation_include_unreferenced_media=False,
        stt_providers=("speechmatics", "deepgram", "openai_compatible"),
        stt_language="fa",
        speechmatics_api_key="sm-key",
        speechmatics_base_url="https://example.test/v2",
        speechmatics_model="enhanced",
        deepgram_api_key="dg-key",
        deepgram_base_url="https://api.deepgram.com/v1",
        deepgram_model="nova-3",
        stt_openai_base_url="",
        stt_openai_api_key="",
        stt_openai_model="whisper-1",
        stt_openai_max_upload=25_000_000,
        stt_transfer_timeout=600.0,
        max_stt_response_bytes=64 * 1024 * 1024,
        stt_status_long_poll=25.0,
        note_providers=("openai_compatible", "nara", "openrouter"),
        note_provider="openai_compatible",
        note_api_key="note-key",
        note_api_base_url="https://note.example/v1",
        note_model="test-model",
        nara_api_key="",
        nara_base_url="https://router.bynara.id/v1",
        nara_model="agnes-2.5-flash",
        openrouter_api_key="",
        openrouter_base_url="https://openrouter.ai/api/v1",
        openrouter_model="google/gemma-4-31b-it:free",
        groq_api_key="",
        groq_base_url=DEFAULT_GROQ_BASE_URL,
        groq_model=DEFAULT_GROQ_MODEL,
        cerebras_api_key="",
        cerebras_base_url=DEFAULT_CEREBRAS_BASE_URL,
        cerebras_model=DEFAULT_CEREBRAS_MODEL,
        gemini_api_key="",
        gemini_model=DEFAULT_GEMINI_MODEL,
        # Zero by default, matching the shipped configuration: note requests
        # must never wait on hidden reasoning tokens.
        gemini_thinking_budget=DEFAULT_GEMINI_THINKING_BUDGET,
        note_context_tokens=131072,
        note_mode="full",
        note_max_output_tokens=4096,
        note_request_timeout=300.0,
        note_json_mode=False,
        note_chunk_chars=64000,
        note_concurrency=2,
        note_compile=True,
        note_compile_max_chars=0,
        note_reasoning_effort="",
        note_strict_grounding=False,
        note_quality_gate="repair",
        note_openrouter_throughput=False,
        note_structured_output=False,
        note_progressive_booklet=False,
        # Pinned to the pre-change behaviour on purpose: the shipped defaults
        # (25 s first-byte guard, 60 s hedge) are covered by tests that build
        # settings from the environment instead of this fixture.
        note_ttfb_timeout=0.0,
        note_hedge_after=0.0,
        note_compile_provider="",
        raw_text_profile="plain",
        raw_fallback_file_chars=100_000,
        raw_fallback_max_files=8,
        office_enabled=True,
        office_max_input_bytes=1_000_000,
        office_max_chars=20_000,
        office_max_slides=30,
        office_session_ttl_minutes=30,
        office_min_interval_seconds=0.0,
        payment_card_number="5022291332906625",
        payment_card_holder="تستر",
        payment_bank_name="بانک تست",
        receipt_dir=Path("receipts"),
        receipt_retention_days=90,
        max_payment_receipt_bytes=5_000_000,
        transcript_retention_days=0,
        free_plan_hours=1,
        plan_5_hours=5,
        plan_5_price_toman=50_000,
        plan_5_validity_days=30,
        plan_10_hours=10,
        plan_10_price_toman=75_000,
        plan_10_validity_days=30,
        plan_20_hours=20,
        plan_20_price_toman=130_000,
        plan_20_validity_days=30,
        plan_25_hours=25,
        plan_25_price_toman=150_000,
        plan_25_validity_days=30,
        plan_50_hours=50,
        plan_50_price_toman=250_000,
        plan_50_validity_days=30,
        max_concurrent_downloads=2,
        max_concurrent_uploads=2,
        stt_fallback_enabled=True,
        speechmatics_max_in_flight=2,
        deepgram_max_in_flight=4,
        stt_openai_max_in_flight=2,
        stt_provider_timeout=1800.0,
        stt_slot_wait_timeout=15.0,
        llm_primary_max_in_flight=4,
        llm_fallback_max_in_flight=4,
        llm_slot_wait_timeout=30.0,
        dedup_window_seconds=120.0,
        plan_values={},
    )
    return replace(base, **overrides) if overrides else base


def wav_bytes(
    duration: float, *, rate: int = 16000, channels: int = 1
) -> bytes:
    """A real PCM WAV file of exactly ``duration`` seconds (digital silence)."""
    frames = int(round(duration * rate))
    buffer = io.BytesIO()
    with wave.open(buffer, "wb") as audio:
        audio.setnchannels(channels)
        audio.setsampwidth(2)
        audio.setframerate(rate)
        audio.writeframes(b"\0\0" * frames * channels)
    return buffer.getvalue()


def video_bytes(
    *, duration: float = 1.0, with_audio: bool = True, rate: int = 15
) -> bytes:
    """A real MP4 (H.264, optionally AAC) built with the bundled encoders."""
    import av

    from fractions import Fraction

    width, height = 64, 48
    with tempfile.TemporaryDirectory() as folder:
        target = Path(folder) / "clip.mp4"
        with av.open(str(target), "w") as container:
            video = container.add_stream("libx264", rate=rate)
            video.width = width
            video.height = height
            video.pix_fmt = "yuv420p"
            audio = None
            if with_audio:
                audio = container.add_stream("aac", rate=44100)
                audio.layout = "stereo"
            total = max(1, int(round(duration * rate)))
            for index in range(total):
                frame = av.VideoFrame(width, height, "yuv420p")
                for plane in frame.planes:
                    view = memoryview(plane)
                    view[: len(view)] = bytes([(index * 8) % 255]) * len(view)
                frame.pts = index
                frame.time_base = Fraction(1, rate)
                for packet in video.encode(frame):
                    container.mux(packet)
                if audio is not None:
                    samples = 44100 // rate
                    audio_frame = av.AudioFrame(
                        format="s16", layout="stereo", samples=samples
                    )
                    audio_frame.sample_rate = 44100
                    audio_frame.time_base = Fraction(1, 44100)
                    audio_frame.pts = index * samples
                    for plane in audio_frame.planes:
                        view = memoryview(plane)
                        view[: len(view)] = b"\0" * len(view)
                    for packet in audio.encode(audio_frame):
                        container.mux(packet)
            for packet in video.encode(None):
                container.mux(packet)
            if audio is not None:
                for packet in audio.encode(None):
                    container.mux(packet)
        return target.read_bytes()


def webm_bytes(
    *, duration: float = 1.0, with_audio: bool = True, rate: int = 15
) -> bytes:
    """A real WebM screen recording (VP9, optionally Opus) built with PyAV."""
    import av

    from fractions import Fraction
    from av.audio.frame import AudioFrame

    width, height = 64, 48
    with tempfile.TemporaryDirectory() as folder:
        target = Path(folder) / "clip.webm"
        with av.open(str(target), "w", format="webm") as container:
            video = container.add_stream("libvpx-vp9", rate=rate)
            video.width = width
            video.height = height
            video.pix_fmt = "yuv420p"
            audio = container.add_stream("libopus", rate=48000) if with_audio else None
            if audio is not None:
                audio.layout = "mono"

            total = max(1, int(round(duration * rate)))
            samples = 48000 // rate
            for index in range(total):
                frame = av.VideoFrame(width, height, "yuv420p")
                for plane, value in zip(frame.planes, (32 + index % 16, 128, 128), strict=True):
                    view = memoryview(plane)
                    view[: len(view)] = bytes([value]) * len(view)
                frame.pts = index
                frame.time_base = Fraction(1, rate)
                for packet in video.encode(frame):
                    container.mux(packet)

                if audio is not None:
                    audio_frame = AudioFrame(
                        format="s16", layout="mono", samples=samples
                    )
                    audio_frame.sample_rate = 48000
                    audio_frame.time_base = Fraction(1, 48000)
                    audio_frame.pts = index * samples
                    for plane in audio_frame.planes:
                        view = memoryview(plane)
                        view[: len(view)] = b"\0" * len(view)
                    for packet in audio.encode(audio_frame):
                        container.mux(packet)

            for packet in video.encode(None):
                container.mux(packet)
            if audio is not None:
                for packet in audio.encode(None):
                    container.mux(packet)
        return target.read_bytes()


def mp3_bytes(duration: float, *, rate: int = 16000) -> bytes:
    """A real MP3 file produced by the bundled libmp3lame encoder."""
    import av

    from fractions import Fraction

    with tempfile.TemporaryDirectory() as folder:
        target = Path(folder) / "audio.mp3"
        with av.open(str(target), "w") as container:
            stream = container.add_stream("mp3", rate=rate)
            stream.layout = "mono"
            chunk = rate // 10
            total = int(round(duration * rate))
            sent = 0
            while sent < total:
                samples = min(chunk, total - sent)
                frame = av.AudioFrame(format="s16", layout="mono", samples=samples)
                frame.sample_rate = rate
                frame.time_base = Fraction(1, rate)
                frame.pts = sent
                for plane in frame.planes:
                    view = memoryview(plane)
                    view[: len(view)] = b"\0" * len(view)
                for packet in stream.encode(frame):
                    container.mux(packet)
                sent += samples
            for packet in stream.encode(None):
                container.mux(packet)
        return target.read_bytes()


def build_deck(
    path: Path,
    slides: list[dict] | None = None,
    orphan_media: dict[str, bytes] | None = None,
) -> Path:
    """Create a .pptx whose slides carry real text plus injected media relations.

    ``slides`` entries accept ``title``, ``bullets``, ``notes``, ``media``
    (a list of ``(part_filename, payload, rel_type)`` tuples), ``chart``
    (``{"categories": [...], "series": [(name, values), ...]}``) and
    ``smartart`` (a list of node texts, written to a real diagram-data part).
    """
    from pptx import Presentation

    slides = slides or []
    deck = Presentation()
    layout = deck.slide_layouts[1]
    for spec in slides:
        slide = deck.slides.add_slide(layout)
        slide.shapes.title.text = spec.get("title", "")
        body = slide.placeholders[1].text_frame
        bullets = spec.get("bullets") or []
        if bullets:
            body.text = bullets[0]
            for extra in bullets[1:]:
                body.add_paragraph().text = extra
        chart = spec.get("chart")
        if chart:
            from pptx.chart.data import CategoryChartData
            from pptx.enum.chart import XL_CHART_TYPE
            from pptx.util import Inches

            data = CategoryChartData()
            data.categories = chart.get("categories") or ()
            for name, values in chart.get("series") or ():
                data.add_series(name, values)
            slide.shapes.add_chart(
                XL_CHART_TYPE.COLUMN_CLUSTERED,
                Inches(1), Inches(3), Inches(5), Inches(3), data,
            )
        if spec.get("notes"):
            slide.notes_slide.notes_text_frame.text = spec["notes"]
    deck.save(str(path))

    content_types = {
        "m4a": "audio/mp4",
        "mp3": "audio/mpeg",
        "wav": "audio/wav",
        "ogg": "audio/ogg",
        "mp4": "video/mp4",
        "mov": "video/quicktime",
    }
    members: dict[str, bytes] = {}
    with zipfile.ZipFile(path) as archive:
        order = archive.namelist()
        for name in order:
            members[name] = archive.read(name)

    for index, spec in enumerate(slides, start=1):
        media = spec.get("media") or []
        if not media:
            continue
        rels_name = f"ppt/slides/_rels/slide{index}.xml.rels"
        rels = members[rels_name].decode("utf-8")
        additions = []
        for position, (part_filename, payload, rel_type) in enumerate(media, start=90):
            members[f"ppt/media/{part_filename}"] = payload
            additions.append(
                f'<Relationship Id="rIdMedia{index}_{position}" Type="{rel_type}" '
                f'Target="../media/{part_filename}"/>'
            )
        rels = rels.replace("</Relationships>", "".join(additions) + "</Relationships>")
        members[rels_name] = rels.encode("utf-8")

    for index, spec in enumerate(slides, start=1):
        points = spec.get("smartart") or []
        if not points:
            continue
        # SmartArt: the words live in ppt/diagrams/dataN.xml and the slide only
        # holds a graphicFrame pointing at it — exactly what python-pptx cannot
        # read, so the fixture has to be assembled at the package level.
        part = f"ppt/diagrams/data{index}.xml"
        members[part] = (
            '<dgm:dataModel '
            'xmlns:dgm="http://schemas.openxmlformats.org/drawingml/2006/diagram" '
            'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
            "<dgm:ptLst>"
            '<dgm:pt modelId="0" type="doc"><dgm:t><a:p>'
            "<a:endParaRPr/></a:p></dgm:t></dgm:pt>"
            + "".join(
                f'<dgm:pt modelId="{position + 1}"><dgm:t><a:p>'
                f"<a:r><a:t>{point}</a:t></a:r></a:p></dgm:t></dgm:pt>"
                for position, point in enumerate(points)
            )
            + "</dgm:ptLst></dgm:dataModel>"
        ).encode("utf-8")
        rels_name = f"ppt/slides/_rels/slide{index}.xml.rels"
        rels = members[rels_name].decode("utf-8")
        rels = rels.replace(
            "</Relationships>",
            f'<Relationship Id="rIdDiagram{index}" '
            'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/diagramData" '
            f'Target="../diagrams/data{index}.xml"/></Relationships>',
        )
        members[rels_name] = rels.encode("utf-8")
        slide_name = f"ppt/slides/slide{index}.xml"
        slide_xml = members[slide_name].decode("utf-8")
        frame = (
            "<p:graphicFrame>"
            '<p:nvGraphicFramePr><p:cNvPr id="900" name="SmartArt"/>'
            "<p:cNvGraphicFramePr/><p:nvPr/></p:nvGraphicFramePr>"
            '<p:xfrm><a:off x="0" y="0"/><a:ext cx="914400" cy="914400"/></p:xfrm>'
            '<a:graphic><a:graphicData '
            'uri="http://schemas.openxmlformats.org/drawingml/2006/diagram">'
            '<dgm:relIds '
            'xmlns:dgm="http://schemas.openxmlformats.org/drawingml/2006/diagram" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
            f'r:dm="rIdDiagram{index}"/>'
            "</a:graphicData></a:graphic></p:graphicFrame>"
        )
        slide_xml = slide_xml.replace("</p:spTree>", frame + "</p:spTree>", 1)
        members[slide_name] = slide_xml.encode("utf-8")

    for name, payload in (orphan_media or {}).items():
        members[f"ppt/media/{name}"] = payload

    # OPC requires a content type for every media extension used in the package.
    declared = members["[Content_Types].xml"].decode("utf-8")
    extensions = {
        Path(name).suffix.lstrip(".").lower()
        for name in members
        if name.startswith("ppt/media/")
    }
    defaults = "".join(
        f'<Default Extension="{extension}" ContentType="{content_types[extension]}"/>'
        for extension in sorted(extensions)
        if extension in content_types and f'Extension="{extension}"' not in declared
    )
    members["[Content_Types].xml"] = declared.replace("<Override", defaults + "<Override", 1).encode("utf-8")

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, payload in members.items():
            archive.writestr(name, payload)
    return path
