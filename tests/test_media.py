"""Real (unfaked) tests for the media layer: probing, extraction and merging.

Every other media test fakes ``probe_media`` / ``extract_audio_track`` at the
bot seam, so the code that actually decodes audio — ``media_worker.py`` and the
PowerPoint reader in ``presentations.py`` — had no direct coverage. These tests
drive the real worker subprocess and PyAV against files built by
``tests.support`` (WAV/MP4/WebM/MP3) and the committed legacy ``.ppt`` decks,
so a regression in extraction from a video or a deck is caught here.
"""

from __future__ import annotations

import asyncio
import functools
import io
import json
import sys
import tempfile
import unittest
import zipfile
import struct
import zlib
from pathlib import Path
from unittest import mock

from gamas_bot import media_worker
from gamas_bot.bot import GamasBot
from gamas_bot.media import (
    AudioMergeResult,
    MediaInfo,
    MergedClipSpan,
    MediaToolError,
    check_media_worker,
    choose_merge_format,
    convert_to_pptx,
    extract_audio_track,
    merge_audio_tracks,
    probe_media,
    probe_media_many,
    parse_merge_output,
)
from gamas_bot.presentations import (
    ClipSpan,
    MediaClip,
    PreparedAudio,
    PresentationContent,
    PresentationError,
    SlideText,
    _extract_slide_text_from_xml,
    _extract_slide_text_with_pptx,
    _preflight_zip_directory,
    _read_member,
    _validate_archive,
    MAX_ZIP_CENTRAL_DIRECTORY_BYTES,
    align_narration_to_slides,
    prepare_audio,
    read_presentation,
    slides_outline,
)
from gamas_bot.transcript import RawTranscript, Word
from tests.support import (
    AUDIO_REL,
    FIXTURES_DIR,
    VIDEO_REL,
    build_deck,
    make_settings,
    mp3_bytes,
    video_bytes,
    wav_bytes,
)

# Generated once per process: encoding a clip is the slow part of these tests.
_wav = functools.lru_cache(maxsize=None)(
    lambda duration: wav_bytes(duration=duration)
)
_video = functools.lru_cache(maxsize=None)(
    lambda duration, with_audio: video_bytes(duration=duration, with_audio=with_audio)
)
_mp3 = functools.lru_cache(maxsize=None)(lambda duration: mp3_bytes(duration))


class _AudioPipelineStub:
    """Just enough surface for ``GamasBot._prepare_audio`` (it reads settings only)."""

    def __init__(self, settings):
        self.settings = settings


class MediaTempCase(unittest.IsolatedAsyncioTestCase):
    def setUp(self) -> None:
        self.settings = make_settings()
        self._folder = tempfile.TemporaryDirectory()
        self.tmp = Path(self._folder.name)

    def tearDown(self) -> None:
        self._folder.cleanup()

    def write(self, name: str, payload: bytes) -> Path:
        path = self.tmp / name
        path.write_bytes(payload)
        return path


class ProbeTests(MediaTempCase):
    async def test_probe_reports_duration_and_codec_for_audio(self):
        path = self.write("a.wav", _wav(1.0))
        info = await probe_media(path, self.settings)
        self.assertTrue(info.has_audio)
        self.assertEqual(info.audio_codec, "pcm_s16le")
        self.assertIsNotNone(info.duration)
        self.assertAlmostEqual(info.duration, 1.0, delta=0.2)

    async def test_probe_detects_a_video_audio_track(self):
        path = self.write("v.mp4", _video(1.0, True))
        info = await probe_media(path, self.settings)
        self.assertTrue(info.has_audio)
        self.assertEqual(info.audio_codec, "aac")

    async def test_probe_reports_no_audio_for_a_silent_video(self):
        path = self.write("silent.mp4", _video(1.0, False))
        info = await probe_media(path, self.settings)
        self.assertFalse(info.has_audio)
        self.assertIsNone(info.audio_codec)

    async def test_batched_probe_matches_the_reported_files(self):
        audio = self.write("a.wav", _wav(1.0))
        video = self.write("v.mp4", _video(1.0, True))
        silent = self.write("silent.mp4", _video(1.0, False))
        reports = await probe_media_many([audio, video, silent], self.settings)
        self.assertEqual({str(audio), str(video), str(silent)}, set(reports))
        self.assertTrue(reports[str(audio)].has_audio)
        self.assertTrue(reports[str(video)].has_audio)
        self.assertFalse(reports[str(silent)].has_audio)


class ExtractTests(MediaTempCase):
    async def test_extract_pulls_a_mono_wav_track_from_video(self):
        source = self.write("v.mp4", _video(1.0, True))
        out = await extract_audio_track(source, self.tmp / "out", self.settings)
        self.assertEqual(out.suffix, ".wav")
        info = await probe_media(out, self.settings)
        self.assertTrue(info.has_audio)
        self.assertEqual(info.audio_codec, "pcm_s16le")
        self.assertAlmostEqual(info.duration, 1.0, delta=0.3)

    async def test_extract_uses_opus_when_wav_would_exceed_the_limit(self):
        source = self.write("v.mp4", _video(1.0, True))
        settings = make_settings(presentation_wav_limit_bytes=1024)
        out = await extract_audio_track(
            source, self.tmp / "out", settings, total_duration=600.0
        )
        self.assertEqual(out.suffix, ".ogg")
        info = await probe_media(out, settings)
        self.assertTrue(info.has_audio)
        self.assertEqual(info.audio_codec, "opus")

    async def test_merge_concatenates_with_padding_silence(self):
        first = self.write("a.wav", _wav(1.0))
        second = self.write("b.wav", _wav(1.0))
        settings = make_settings(presentation_silence_seconds=0.5)
        result = await merge_audio_tracks(
            [first, second], self.tmp / "merged", settings,
            total_duration=2.5, stem="m",
        )
        info = await probe_media(result.path, self.settings)
        self.assertTrue(info.has_audio)
        self.assertAlmostEqual(info.duration, 2.5, delta=0.3)
        self.assertAlmostEqual(result.source_duration, 2.0, delta=0.05)
        self.assertAlmostEqual(result.output_duration, 2.5, delta=0.05)
        self.assertEqual(len(result.clip_spans), 2)
        self.assertAlmostEqual(result.clip_spans[0].start_seconds, 0.0, delta=0.01)
        self.assertAlmostEqual(result.clip_spans[0].end_seconds, 1.0, delta=0.05)
        self.assertAlmostEqual(result.clip_spans[1].start_seconds, 1.5, delta=0.05)
        self.assertAlmostEqual(result.clip_spans[1].end_seconds, 2.5, delta=0.05)


class MediaOperationLimitTests(unittest.IsolatedAsyncioTestCase):
    async def test_media_stages_obey_the_configured_semaphore(self):
        stub = _AudioPipelineStub(make_settings(media_worker_concurrency=1))
        stub._media_semaphore = asyncio.Semaphore(stub.settings.media_worker_concurrency)
        active = 0
        high_water = 0

        async def operation():
            nonlocal active, high_water
            active += 1
            high_water = max(high_water, active)
            await asyncio.sleep(0.005)
            active -= 1

        await asyncio.gather(*(
            GamasBot._run_media_operation(stub, operation)
            for _ in range(4)
        ))
        self.assertEqual(high_water, 1)


class MergeReportTests(unittest.TestCase):
    def test_rejects_a_merge_report_with_inconsistent_clip_accounting(self):
        payload = json.dumps({
            "sample_rate": 16000,
            "source_samples": 32000,
            "output_samples": 36000,
            "clip_spans": [
                {"index": 0, "start_sample": 0, "end_sample": 16000},
                {"index": 1, "start_sample": 20000, "end_sample": 30000},
            ],
        })
        with self.assertRaises(MediaToolError):
            parse_merge_output(payload, 2)


class FormatChoiceTests(unittest.TestCase):
    def test_choose_merge_format_switches_to_opus_above_the_limit(self):
        settings = make_settings(presentation_wav_limit_bytes=32000)
        self.assertEqual(choose_merge_format(None, settings), "wav")
        self.assertEqual(choose_merge_format(0.5, settings), "wav")
        self.assertEqual(choose_merge_format(10.0, settings), "opus")

    def test_merge_format_respects_remaining_temp_and_file_budget(self):
        settings = make_settings(presentation_wav_limit_bytes=1_000_000)
        self.assertEqual(
            choose_merge_format(10.0, settings, max_output_bytes=200_000), "opus"
        )
        with self.assertRaises(MediaToolError):
            choose_merge_format(10.0, settings, max_output_bytes=100_000)


class FormatPolicyTests(unittest.TestCase):
    """``AUDIO_FORMAT_POLICY`` decides by time once a transfer rate is known."""

    def test_auto_prefers_opus_when_the_upload_is_the_bottleneck(self):
        settings = make_settings(audio_format_policy="auto")
        # 2 h: 230 MB of PCM against 29 MB of Opus. At 2 MB/s the PCM upload
        # costs ~115 s more than the whole Opus encode.
        self.assertEqual(
            choose_merge_format(7200.0, settings, egress_bytes_per_second=2e6),
            "opus",
        )

    def test_auto_keeps_pcm_when_the_link_is_faster_than_the_encoder(self):
        settings = make_settings(audio_format_policy="auto")
        # At 100 MB/s the transfer is irrelevant and PCM's cheaper encode wins.
        self.assertEqual(
            choose_merge_format(7200.0, settings, egress_bytes_per_second=100e6),
            "wav",
        )

    def test_auto_without_a_measured_rate_keeps_the_size_only_rule(self):
        settings = make_settings(
            audio_format_policy="auto", presentation_wav_limit_bytes=1000
        )
        self.assertEqual(choose_merge_format(10.0, settings), "opus")

    def test_explicit_policies_override_the_costing(self):
        fast = make_settings(audio_format_policy="wav")
        slow = make_settings(audio_format_policy="opus")
        self.assertEqual(
            choose_merge_format(7200.0, fast, egress_bytes_per_second=1e6), "wav"
        )
        self.assertEqual(
            choose_merge_format(7200.0, slow, egress_bytes_per_second=1e6), "opus"
        )

    def test_a_forced_pcm_policy_still_respects_the_budget(self):
        settings = make_settings(
            audio_format_policy="wav", presentation_wav_limit_bytes=32_000
        )
        # PCM does not fit the budget, so the smaller container is produced
        # rather than an upload no provider could accept.
        self.assertEqual(choose_merge_format(3600.0, settings), "opus")

    def test_neither_container_fitting_is_an_error(self):
        settings = make_settings(audio_format_policy="auto")
        with self.assertRaises(MediaToolError):
            choose_merge_format(
                3600.0, settings, max_output_bytes=1_000, egress_bytes_per_second=1e6
            )


class PrepareAudioRoutingTests(MediaTempCase):
    """The bot must forward the measured duration to extraction."""

    async def test_prepare_audio_forwards_the_probed_duration(self):
        captured: dict = {}

        async def fake_probe(_path, _settings):
            return MediaInfo(True, 123.5, "aac")

        async def fake_extract(
            source, _workdir, _settings, *, total_duration=None,
            max_output_bytes=None, egress_bytes_per_second=None, stem="audio",
        ):
            captured["total_duration"] = total_duration
            captured["max_output_bytes"] = max_output_bytes
            captured["egress_bytes_per_second"] = egress_bytes_per_second
            return source

        stub = _AudioPipelineStub(self.settings)
        with mock.patch("gamas_bot.bot.probe_media", new=fake_probe), mock.patch(
            "gamas_bot.bot.extract_audio_track", new=fake_extract
        ):
            result = await GamasBot._prepare_audio(
                stub,
                Path("v.mp4"),
                self.tmp,
                "v.mp4",
                "video",
                egress_bytes_per_second=2_000_000.0,
            )
        self.assertEqual(captured["total_duration"], 123.5)
        # The prepared file is sized to the STT chain's largest ceiling, not
        # only to MAX_FILE_SIZE: a WAV no provider accepts is never produced.
        self.assertIsNotNone(captured["max_output_bytes"])
        self.assertLessEqual(captured["max_output_bytes"], self.settings.max_file_size)
        # The measured transfer rate reaches the format decision, which is what
        # makes AUDIO_FORMAT_POLICY=auto a time trade rather than a size trade.
        self.assertEqual(captured["egress_bytes_per_second"], 2_000_000.0)
        self.assertEqual(result, Path("v.mp4"))

    async def test_passthrough_audio_is_not_transcoded(self):
        async def fake_probe(_path, _settings):
            return MediaInfo(True, 10.0, "mp3")

        async def fail_extract(*_args, **_kwargs):  # pragma: no cover - must not run
            raise AssertionError("passthrough audio must not be extracted")

        stub = _AudioPipelineStub(self.settings)
        source = self.write("a.mp3", _mp3(0.5))
        with mock.patch("gamas_bot.bot.probe_media", new=fake_probe), mock.patch(
            "gamas_bot.bot.extract_audio_track", new=fail_extract
        ):
            result = await GamasBot._prepare_audio(
                stub, source, self.tmp, "a.mp3", "audio"
            )
        self.assertEqual(result, source)


class PresentationTests(MediaTempCase):
    def deck_with_media(self) -> Path:
        return build_deck(
            self.tmp / "deck.pptx",
            slides=[
                {"title": "One", "notes": "n", "media": [("a.wav", _wav(1.5), AUDIO_REL)]},
                {"title": "Two", "media": [("v.mp4", _video(1.0, True), VIDEO_REL)]},
                {"title": "Three", "media": [("silent.mp4", _video(1.0, False), VIDEO_REL)]},
            ],
            orphan_media={"orphan.mp3": _mp3(0.8)},
        )

    async def test_deck_extracts_and_merges_audio_and_video_narration(self):
        deck = self.deck_with_media()
        settings = make_settings(presentation_include_unreferenced_media=True)
        content = read_presentation(deck, self.tmp / "media", settings)
        self.assertEqual(
            [(clip.slide_number, clip.kind) for clip in content.clips],
            [(1, "audio"), (2, "video"), (3, "video"), (None, "audio")],
        )
        prepared = await prepare_audio(content, self.tmp / "work", settings)
        self.assertIsNotNone(prepared)
        # The silent video is dropped; the other three clips are merged.
        self.assertEqual(len(prepared.clips), 3)
        self.assertTrue(
            any("بدون شاخهٔ صوتی" in reason for reason in prepared.skipped)
        )
        info = await probe_media(prepared.path, self.settings)
        self.assertTrue(info.has_audio)
        # The measured audio sum excludes two synthetic 0.25 s gaps.
        self.assertAlmostEqual(info.duration, prepared.output_duration, delta=0.03)
        self.assertAlmostEqual(
            prepared.output_duration, prepared.source_duration + 0.5, delta=0.01
        )
        self.assertEqual([span.clip.slide_number for span in prepared.clip_spans], [1, 2, None])
        self.assertAlmostEqual(prepared.clip_spans[1].start_seconds, 1.75, delta=0.1)
        self.assertAlmostEqual(
            prepared.clip_spans[2].start_seconds,
            prepared.clip_spans[1].end_seconds + 0.25,
            delta=0.01,
        )

    async def test_one_compatible_audio_clip_uses_the_extracted_source_directly(self):
        deck = build_deck(
            self.tmp / "single.pptx",
            slides=[{
                "title": "One",
                "media": [("voice.wav", _wav(0.8), AUDIO_REL)],
            }],
        )
        content = read_presentation(deck, self.tmp / "media", self.settings)
        prepared = await prepare_audio(content, self.tmp / "work", self.settings)
        self.assertIsNotNone(prepared)
        assert prepared is not None
        self.assertFalse(prepared.merged)
        self.assertEqual(prepared.path, content.clips[0].path)
        self.assertAlmostEqual(prepared.source_duration, 0.8, delta=0.1)
        self.assertAlmostEqual(prepared.output_duration, prepared.source_duration, delta=0.01)
        self.assertEqual(len(prepared.clip_spans), 1)
        self.assertEqual(prepared.clip_spans[0].clip.slide_number, 1)

    def test_unreferenced_media_is_excluded_by_default(self):
        deck = self.deck_with_media()
        content = read_presentation(deck, self.tmp / "media", self.settings)
        self.assertEqual([clip.slide_number for clip in content.clips], [1, 2, 3])
        self.assertTrue(
            any("بدون پیوند مستقیم" in reason for reason in content.skipped),
            content.skipped,
        )

    def test_reused_media_part_keeps_each_slide_reference(self):
        payload = _wav(0.8)
        deck = build_deck(
            self.tmp / "reused.pptx",
            slides=[
                {"title": "One", "media": [("shared.wav", payload, AUDIO_REL)]},
                {"title": "Two", "media": [("shared.wav", payload, AUDIO_REL)]},
            ],
        )
        content = read_presentation(deck, self.tmp / "media", self.settings)
        self.assertEqual([clip.slide_number for clip in content.clips], [1, 2])
        self.assertEqual(content.clips[0].path, content.clips[1].path)
        self.assertEqual([clip.order for clip in content.clips], [1, 2])

    async def test_video_audio_is_dropped_when_disabled(self):
        deck = self.deck_with_media()
        settings = make_settings(presentation_include_video_audio=False)
        content = read_presentation(deck, self.tmp / "media", settings)
        # Video clips are dropped; plain audio narration is still kept.
        self.assertFalse(any(clip.kind == "video" for clip in content.clips))
        self.assertTrue(any(clip.kind == "audio" for clip in content.clips))
        self.assertTrue(
            any("ویدیو نادیده گرفته شد" in reason for reason in content.skipped)
        )
        self.assertIsNotNone(await prepare_audio(content, self.tmp / "work", settings))

    async def test_too_short_clip_is_skipped_not_merged(self):
        deck = build_deck(
            self.tmp / "short.pptx",
            slides=[{"title": "S", "media": [("s.wav", _wav(0.2), AUDIO_REL)]}],
        )
        content = read_presentation(deck, self.tmp / "media", self.settings)
        skipped: list[str] = []
        prepared = await prepare_audio(
            content, self.tmp / "work", self.settings, skipped_out=skipped
        )
        self.assertIsNone(prepared)
        # The Persian reason carries ZWNJ; the file name is the stable anchor.
        self.assertEqual(len(skipped), 1, skipped)
        self.assertIn("s.wav", skipped[0])

    async def test_text_only_deck_has_no_audio_but_keeps_slides(self):
        deck = build_deck(
            self.tmp / "text.pptx",
            slides=[{"title": "Only", "bullets": ["x"]}],
        )
        content = read_presentation(deck, self.tmp / "media", self.settings)
        self.assertEqual(content.clips, ())
        self.assertTrue(content.slides)
        self.assertIsNone(await prepare_audio(content, self.tmp / "work", self.settings))

    def test_outline_contains_slide_text_and_speaker_notes(self):
        deck = build_deck(
            self.tmp / "notes.pptx",
            slides=[{
                "title": "موضوع ارائه",
                "bullets": ["نکتهٔ روی اسلاید"],
                "notes": "توضیح تکمیلی گوینده",
            }],
        )
        content = read_presentation(deck, self.tmp / "media", self.settings)
        outline = slides_outline(content.slides)
        self.assertIn("موضوع ارائه", outline)
        self.assertIn("نکتهٔ روی اسلاید", outline)
        self.assertIn("یادداشت گوینده: توضیح تکمیلی گوینده", outline)


class ArchiveValidationTests(MediaTempCase):
    def test_part_names_with_control_characters_are_rejected(self):
        """A newline in an entry name would let a deck forge log lines (F-19)."""
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr("ppt/slides/slide1.xml\nFAKE", "x")
        buffer.seek(0)
        with zipfile.ZipFile(buffer) as archive:
            with self.assertRaises(PresentationError):
                _validate_archive(archive, self.settings)

    def test_overlong_member_names_are_rejected(self):
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr("a" * 4097, "x")
        buffer.seek(0)
        with zipfile.ZipFile(buffer) as archive:
            with self.assertRaises(PresentationError):
                _validate_archive(archive, self.settings)

    def test_oversized_central_directory_is_rejected_before_zipfile_indexing(self):
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr("ppt/slides/slide1.xml", "x")
        forged = bytearray(buffer.getvalue())
        end_record = forged.rfind(b"PK\x05\x06")
        self.assertGreaterEqual(end_record, 0)
        struct.pack_into(
            "<I",
            forged,
            end_record + 12,
            MAX_ZIP_CENTRAL_DIRECTORY_BYTES + 1,
        )
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "large-index.zip"
            path.write_bytes(forged)
            with self.assertRaises(PresentationError):
                _preflight_zip_directory(path)

    def test_declared_zip_size_cannot_hide_actual_member_data(self):
        payload = b"actual expanded member contents" * 100
        for compression in (zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED):
            with self.subTest(compression=compression):
                buffer = io.BytesIO()
                with zipfile.ZipFile(buffer, "w", compression=compression) as archive:
                    archive.writestr("ppt/media/clip.bin", payload)
                forged = bytearray(buffer.getvalue())
                central = forged.find(b"PK\x01\x02")
                self.assertGreaterEqual(central, 0)
                # Forge a one-byte uncompressed size and matching one-byte CRC,
                # while retaining the original compressed length and contents.
                struct.pack_into("<I", forged, central + 16, zlib.crc32(payload[:1]))
                struct.pack_into("<I", forged, central + 24, 1)
                with zipfile.ZipFile(io.BytesIO(forged)) as archive:
                    with self.assertRaises(PresentationError):
                        _read_member(archive, "ppt/media/clip.bin", len(payload) + 1)


class PresentationAlignmentTests(MediaTempCase):
    def prepared(self) -> PreparedAudio:
        clips = (
            MediaClip(1, 1, "ppt/media/one.wav", self.tmp / "one.wav", "audio", 1.0),
            MediaClip(2, 2, "ppt/media/two.wav", self.tmp / "two.wav", "audio", 1.0),
            MediaClip(3, None, "ppt/media/orphan.wav", self.tmp / "orphan.wav", "audio", 1.0),
        )
        return PreparedAudio(
            path=self.tmp / "merged.wav",
            clips=clips,
            source_duration=3.0,
            output_duration=3.5,
            clip_spans=(
                ClipSpan(clips[0], 0.0, 1.0),
                ClipSpan(clips[1], 1.25, 2.25),
                ClipSpan(clips[2], 2.5, 3.5),
            ),
            merged=True,
        )

    def test_timed_words_map_only_to_their_linked_slide_clip(self):
        raw = RawTranscript(
            words=[
                Word("first", 0.1, 0.3),
                Word("gap", 1.05, 1.15),
                Word("second", 1.4, 1.6),
                Word("crosses", 0.8, 1.4),
                Word("orphan", 2.7, 2.9),
                Word("untimed"),
            ],
            text="first gap second orphan untimed",
            engine="test",
        )
        aligned = align_narration_to_slides(
            raw, self.prepared(), (SlideText(1), SlideText(2))
        )
        self.assertTrue(aligned.aligned)
        self.assertEqual(aligned.slide_text, ((1, "first"), (2, "second")))
        self.assertIn("gap", aligned.unassigned_text)
        self.assertIn("crosses", aligned.unassigned_text)
        self.assertIn("orphan", aligned.unassigned_text)
        self.assertIn("untimed", aligned.unassigned_text)
        self.assertNotIn("orphan", aligned.as_text().split("### گفتار بدون", 1)[0])

    def test_plain_text_stays_explicitly_unaligned(self):
        raw = RawTranscript.from_text("بدون زمان‌بندی")
        aligned = align_narration_to_slides(
            raw, self.prepared(), (SlideText(1), SlideText(2))
        )
        self.assertFalse(aligned.aligned)
        self.assertEqual(aligned.slide_text, ())
        self.assertIn("بدون نگاشت مطمئن", aligned.as_text())
        self.assertIn("بدون زمان‌بندی", aligned.as_text())

    def test_empty_plain_transcript_does_not_create_marker_only_text(self):
        raw = RawTranscript.from_text("")
        aligned = align_narration_to_slides(
            raw, self.prepared(), (SlideText(1), SlideText(2))
        )
        self.assertFalse(aligned.aligned)
        self.assertEqual(aligned.as_text(), "")

    def test_a_word_that_spills_past_the_last_clip_stays_with_it(self):
        """The passthrough span ends at the probed duration; MP3 padding jitters."""
        clips = (
            MediaClip(1, 1, "ppt/media/one.wav", self.tmp / "one.wav", "audio", 1.0),
        )
        prepared = PreparedAudio(
            path=self.tmp / "one.wav",
            clips=clips,
            source_duration=1.0,
            output_duration=1.0,
            clip_spans=(ClipSpan(clips[0], 0.0, 1.0),),
            merged=False,
        )
        raw = RawTranscript(
            words=[Word("پایان", 0.9, 1.03)], text="پایان", engine="test"
        )
        aligned = align_narration_to_slides(raw, prepared, (SlideText(1),))
        self.assertTrue(aligned.aligned)
        self.assertEqual(aligned.slide_text, ((1, "پایان"),))

    def test_boundary_jittered_word_stays_with_its_slide_clip(self):
        raw = RawTranscript(
            words=[
                Word("شروع", 0.1, 0.5),
                # Spills 30 ms into the 250 ms synthetic silence gap after clip 1:
                Word("پایان.", 0.85, 1.03),
                # Starts 20 ms before clip 2's nominal 1.25 s boundary:
                Word("اسلاید", 1.23, 1.5),
            ],
            text="شروع پایان. اسلاید",
            engine="test",
        )
        aligned = align_narration_to_slides(
            raw, self.prepared(), (SlideText(1), SlideText(2))
        )
        self.assertTrue(aligned.aligned)
        self.assertEqual(
            aligned.slide_text,
            ((1, "شروع پایان."), (2, "اسلاید")),
        )
        self.assertEqual(aligned.unassigned_text, "")

    def test_slide_table_renders_as_markdown_pipe_rows_in_outline(self):
        slide = SlideText(
            number=1,
            title="جدول مقایسه",
            body=("نام | سن", "علی | ۲۰"),
            notes="خط اول یادداشت\nخط دوم یادداشت",
        )
        markdown = slide.as_markdown()
        self.assertIn("| نام | سن |", markdown)
        self.assertIn("| علی | ۲۰ |", markdown)
        self.assertIn("یادداشت گوینده: خط اول یادداشت\nخط دوم یادداشت", markdown)


class PresentationLimitTests(MediaTempCase):
    async def test_merged_duration_cap_includes_inter_clip_silence(self):
        first = self.write("first.wav", b"fake")
        second = self.write("second.wav", b"fake")
        clips = (
            MediaClip(1, 1, "ppt/media/first.wav", first, "audio"),
            MediaClip(2, 2, "ppt/media/second.wav", second, "audio"),
        )
        content = PresentationContent((SlideText(1), SlideText(2)), clips)
        settings = make_settings(presentation_max_total_duration=1.44)
        merged_path = self.tmp / "work" / "merged.wav"
        merged_path.parent.mkdir()
        merged_path.write_bytes(b"merged")

        async def fake_probe(paths, _settings):
            return {str(path): MediaInfo(True, 0.6, "pcm_s16le") for path in paths}

        captured: dict = {}

        async def fake_merge(*_args, **kwargs):
            captured.update(kwargs)
            return AudioMergeResult(
                merged_path,
                sample_rate=16000,
                source_duration=1.2,
                output_duration=1.45,
                clip_spans=(
                    MergedClipSpan(0, 0.0, 0.6),
                    MergedClipSpan(1, 0.85, 1.45),
                ),
            )

        with mock.patch("gamas_bot.presentations.probe_media_many", new=fake_probe), mock.patch(
            "gamas_bot.presentations.merge_audio_tracks", new=fake_merge
        ):
            with self.assertRaises(PresentationError):
                await prepare_audio(content, merged_path.parent, settings)
        self.assertAlmostEqual(captured["total_duration"], 1.45, delta=0.01)

    def test_extraction_respects_the_per_presentation_temp_budget(self):
        deck = build_deck(
            self.tmp / "large-temp.pptx",
            slides=[{"title": "S", "media": [("clip.wav", _wav(0.8), AUDIO_REL)]}],
        )
        settings = make_settings(
            presentation_max_temp_bytes=deck.stat().st_size + 100
        )
        with self.assertRaises(PresentationError):
            read_presentation(deck, self.tmp / "media", settings)


class SmartArtAndChartTests(MediaTempCase):
    """Text that is not a shape text frame must still reach the outline.

    SmartArt keeps its words in ``ppt/diagrams/dataN.xml`` and a chart keeps
    its title, series names and category labels in ``ppt/charts/chartN.xml``.
    Neither part is read by python-pptx, and both used to leave the slide
    looking empty — so the note model never saw that slide at all.
    """

    def test_smartart_node_text_is_read(self):
        deck = build_deck(
            self.tmp / "smartart.pptx",
            slides=[
                {
                    "title": "فرایند",
                    "smartart": ["تحلیل نیاز", "طراحی", "پیاده‌سازی"],
                }
            ],
        )
        outline = slides_outline(read_presentation(deck, self.tmp / "media", self.settings).slides)
        self.assertIn("تحلیل نیاز", outline)
        self.assertIn("طراحی", outline)
        self.assertIn("پیاده‌سازی", outline)

    def test_chart_title_series_and_categories_are_read(self):
        deck = build_deck(
            self.tmp / "chart.pptx",
            slides=[
                {
                    "title": "عملکرد",
                    "chart": {
                        "categories": ["فروردین", "اردیبهشت"],
                        "series": [("درآمد", (10, 20))],
                    },
                }
            ],
        )
        outline = slides_outline(read_presentation(deck, self.tmp / "media", self.settings).slides)
        self.assertIn("سری‌های نمودار: درآمد", outline)
        self.assertIn("دسته‌های نمودار: فروردین، اردیبهشت", outline)
        self.assertIn("مقادیر سری «درآمد»: فروردین=10، اردیبهشت=20", outline)

    def test_graphics_text_is_read_by_the_xml_fallback_too(self):
        deck = build_deck(
            self.tmp / "graphics.pptx",
            slides=[
                {
                    "title": "فرایند",
                    "smartart": ["گام یک"],
                    "chart": {"categories": ["الف"], "series": [("سری", (1,))]},
                }
            ],
        )
        with mock.patch(
            "gamas_bot.presentations._extract_slide_text_with_pptx",
            new=lambda _path: [],
        ):
            content = read_presentation(deck, self.tmp / "media", self.settings)
        outline = slides_outline(content.slides)
        self.assertIn("گام یک", outline)
        self.assertIn("سری‌های نمودار: سری", outline)
        self.assertIn("دسته‌های نمودار: الف", outline)


class SlideTextFallbackTests(MediaTempCase):
    """The deck's text must reach the note model even without python-pptx.

    ``presentations._extract_slide_text`` used to return ``[]`` whenever
    python-pptx was missing or could not open a particular deck — the narration
    was transcribed and the slide text silently never reached the LLM. The
    package-XML reader is the fallback these tests pin down.
    """

    def build(self) -> Path:
        return build_deck(
            self.tmp / "lecture.pptx",
            slides=[
                {
                    "title": "شبکهٔ عصبی",
                    "bullets": ["لایهٔ ورودی", "لایهٔ پنهان"],
                    "notes": "توضیح تکمیلی گوینده",
                },
                {"title": "آموزش", "bullets": ["گرادیان کاهشی"]},
            ],
        )

    def test_the_xml_fallback_reads_what_python_pptx_reads(self):
        deck = self.build()
        self.assertEqual(
            _extract_slide_text_from_xml(deck),
            _extract_slide_text_with_pptx(deck),
        )

    def test_the_xml_fallback_reads_table_rows_like_the_primary_reader(self):
        from pptx import Presentation
        from pptx.util import Inches

        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        frame = slide.shapes.add_table(
            2, 2, Inches(1), Inches(1), Inches(4), Inches(1)
        )
        table = frame.table
        table.cell(0, 0).text = "نام"
        table.cell(0, 1).text = "سن"
        table.cell(1, 0).text = "علی"
        table.cell(1, 1).text = "۲۰"
        path = self.tmp / "table.pptx"
        deck.save(path)
        self.assertEqual(
            _extract_slide_text_from_xml(path),
            _extract_slide_text_with_pptx(path),
        )

    def test_xml_fallback_fills_each_partially_empty_slide(self):
        deck = self.build()
        partial = [
            SlideText(1, "شبکهٔ عصبی", ("متن اولیه",)),
            SlideText(2),
        ]
        with mock.patch(
            "gamas_bot.presentations._extract_slide_text_with_pptx",
            new=lambda _path: partial,
        ):
            content = read_presentation(deck, self.tmp / "media", self.settings)
        outline = slides_outline(content.slides)
        self.assertIn("متن اولیه", outline)
        self.assertIn("لایهٔ ورودی", outline)
        self.assertIn("توضیح تکمیلی گوینده", outline)
        self.assertIn("آموزش", outline)
        self.assertIn("گرادیان کاهشی", outline)

    def test_slide_text_survives_a_missing_python_pptx(self):
        deck = self.build()
        with mock.patch.dict(sys.modules, {"pptx": None}):
            content = read_presentation(deck, self.tmp / "media", self.settings)
        outline = slides_outline(content.slides)
        self.assertIn("شبکهٔ عصبی", outline)
        self.assertIn("لایهٔ ورودی", outline)
        self.assertIn("یادداشت گوینده: توضیح تکمیلی گوینده", outline)

    def test_slide_text_survives_a_deck_python_pptx_cannot_open(self):
        deck = self.build()
        with mock.patch(
            "gamas_bot.presentations._extract_slide_text_with_pptx",
            new=lambda _path: [],
        ):
            content = read_presentation(deck, self.tmp / "media", self.settings)
        outline = slides_outline(content.slides)
        self.assertIn("شبکهٔ عصبی", outline)
        self.assertIn("گرادیان کاهشی", outline)

    def test_a_text_free_deck_says_why_only_the_narration_is_used(self):
        deck = build_deck(self.tmp / "pictures.pptx", slides=[{"title": ""}])
        content = read_presentation(deck, self.tmp / "media", self.settings)
        self.assertTrue(
            any("متن قابل‌انتخاب" in reason for reason in content.skipped),
            content.skipped,
        )


class LegacyConversionTests(MediaTempCase):
    async def test_bundled_legacy_deck_converts_and_reads_back(self):
        source = FIXTURES_DIR / "visual_minimal.ppt"
        converted = await convert_to_pptx(source, self.tmp / "converted", self.settings)
        self.assertEqual(converted.suffix, ".pptx")
        self.assertGreater(converted.stat().st_size, 0)
        content = read_presentation(converted, self.tmp / "media", self.settings)
        self.assertGreaterEqual(content.slide_count, 1)

    async def test_odp_is_refused_with_a_clear_error(self):
        source = self.write("deck.odp", b"not really an odp")
        with self.assertRaises(MediaToolError):
            await convert_to_pptx(source, self.tmp / "out", self.settings)


class WorkerDependencyTests(MediaTempCase):
    async def test_worker_check_reports_both_dependencies(self):
        payload = json.loads(await check_media_worker())
        self.assertIn("av", payload)
        self.assertIn("ppt2pptx", payload)

    def test_configure_logging_survives_a_missing_av(self):
        with mock.patch.dict(sys.modules, {"av": None}):
            media_worker._configure_logging()  # must not raise

    def test_require_av_reports_a_friendly_error(self):
        with mock.patch.dict(sys.modules, {"av": None}):
            with self.assertRaises(media_worker.WorkerError) as caught:
                media_worker._require_av()
        self.assertIn("media dependencies are missing", str(caught.exception))

    def test_check_reports_a_friendly_error_when_av_is_missing(self):
        with mock.patch.dict(sys.modules, {"av": None}):
            with self.assertRaises(media_worker.WorkerError):
                media_worker._check()


if __name__ == "__main__":
    unittest.main()
