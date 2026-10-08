"""Equivalence of the optimized DOCX hot paths with the public python-docx API.

``_style_run`` writes ``w:rPr`` directly for fresh runs (a booklet styles
thousands of runs, and the generic setters cost several schema-order tree
scans each), and ``_fill_static_toc`` resolves heading styles from the raw
``w:pStyle`` id instead of resolving ``paragraph.style`` per paragraph.

Both are pure performance changes: these tests prove the produced XML is
identical to what the setter-based / public paths produce, so the DOCX output
cannot drift.
"""
from __future__ import annotations

import re
import unittest

from docx import Document
from docx.shared import RGBColor
from lxml import etree

from gamas_bot.docx_export import _heading_level_resolver, _style_run

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

#: Font values covering every role the build can pass: Persian body face,
#: Latin face and the configured Latin-default (Tahoma, the ``w:cs``-skip case).
FONT_VALUES = ("Vazirmatn", "Aptos", "Tahoma")


def rpr_xml(run) -> bytes:
    element = run._r.find(W + "rPr")
    return b"" if element is None else etree.tostring(element)


def _styled_run(document, *, font, size, bold, color, italic, rtl, precreate_rpr):
    run = document.add_paragraph().add_run("x")
    if precreate_rpr:
        # An already-present (empty) w:rPr forces the generic setter-based
        # branch of _style_run, which is the reference implementation.
        run._r.get_or_add_rPr()
    _style_run(
        run, font=font, size=size, bold=bold, color=color,
        italic=italic, rtl=rtl,
    )
    return run


class RunStylingEquivalenceTests(unittest.TestCase):
    def test_fresh_run_matches_setter_path_for_every_combination(self):
        document = Document()
        for font in FONT_VALUES:
            for rtl in (True, False):
                for bold in (False, True):
                    for italic in (False, True):
                        for color in (None, RGBColor(0x1F, 0x3A, 0x5F)):
                            for size in (11, 10.5, 9):
                                with self.subTest(
                                    font=font, rtl=rtl, bold=bold,
                                    italic=italic, color=color, size=size,
                                ):
                                    kwargs = {
                                        "font": font, "size": size, "bold": bold,
                                        "color": color, "italic": italic, "rtl": rtl,
                                    }
                                    fast = _styled_run(
                                        document, precreate_rpr=False, **kwargs
                                    )
                                    legacy = _styled_run(
                                        document, precreate_rpr=True, **kwargs
                                    )
                                    self.assertEqual(
                                        rpr_xml(fast), rpr_xml(legacy),
                                        "optimized run styling diverged",
                                    )

    def test_underline_appended_after_styling_lands_in_the_same_place(self):
        document = Document()
        for rtl in (True, False):
            fast = _styled_run(
                document, font="Vazirmatn", size=10.5, bold=False,
                color=RGBColor(0, 0, 0), italic=False, rtl=rtl,
                precreate_rpr=False,
            )
            fast.font.underline = True
            legacy = _styled_run(
                document, font="Vazirmatn", size=10.5, bold=False,
                color=RGBColor(0, 0, 0), italic=False, rtl=rtl,
                precreate_rpr=True,
            )
            legacy.font.underline = True
            self.assertEqual(rpr_xml(fast), rpr_xml(legacy))

    def test_run_properties_element_precedes_run_content(self):
        document = Document()
        run = document.add_paragraph().add_run("متن")
        _style_run(run, font="Vazirmatn", size=11, rtl=True)
        children = [child.tag for child in run._r]
        self.assertEqual(children[0], W + "rPr")
        self.assertIn(W + "t", children)


class HeadingLevelResolverTests(unittest.TestCase):
    def test_resolver_matches_the_public_style_name_scan(self):
        document = Document()  # default template ships Heading 1..9
        paragraphs = []
        for index in range(40):
            paragraph = document.add_paragraph(f"p{index}")
            if index % 4 == 0:
                paragraph.style = "Heading 1"
            elif index % 4 == 1:
                paragraph.style = "Heading 3"
            elif index % 4 == 2:
                paragraph.style = "Normal"
            paragraphs.append(paragraph)  # index % 4 == 3: no explicit style
        # A pStyle id the document does not define: python-docx resolves it to
        # the default style, and so must the resolver.
        unknown = document.add_paragraph("unknown")
        unknown._p.style = "NoSuchStyle"
        paragraphs.append(unknown)

        resolver = _heading_level_resolver(document)
        for paragraph in paragraphs:
            match = re.fullmatch(r"Heading ([1-3])", paragraph.style.name or "")
            expected = int(match.group(1)) if match else None
            self.assertEqual(
                resolver(paragraph), expected,
                f"resolver disagreed for style {paragraph.style.name!r}",
            )

    def test_plain_paragraph_resolves_to_no_heading(self):
        document = Document()
        resolver = _heading_level_resolver(document)
        self.assertIsNone(resolver(document.add_paragraph("plain")))


if __name__ == "__main__":  # pragma: no cover
    unittest.main()
