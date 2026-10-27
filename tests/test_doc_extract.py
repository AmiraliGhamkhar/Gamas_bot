"""Document ingestion primitives: extraction per format, chunking, dedupe keys.

The fixtures are generated in-memory with the libraries the project already
depends on (reportlab writes the PDF, python-docx the Word file, python-pptx the
deck), so the suite carries no binary blobs and every format is exercised by a
real file rather than a mock.
"""

from __future__ import annotations

import io
import tempfile
import unittest
from pathlib import Path

from gamas_bot.doc_extract import (
    DOCX_MIME,
    PDF_MIME,
    PPTX_MIME,
    DocChunk,
    DocExtractError,
    DocSection,
    ScannedPdfError,
    chunk_sections,
    chunk_text,
    chunks_from_json,
    chunks_to_json,
    classify_document,
    document_title,
    extract_document,
    looks_scanned,
    preview_text,
    sections_from_slides,
    sha256_file,
)
from gamas_bot.learning import MAX_MODEL_INPUT_CHARS
from gamas_bot.progress import to_persian_digits


def make_pdf_bytes(
    text: str = (
        "Cellular respiration produces ATP in the mitochondria of every cell, "
        "and glycolysis happens in the cytoplasm before the Krebs cycle runs "
        "inside the mitochondrial matrix with oxygen as the final acceptor."
    ),
) -> bytes:
    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    chart = canvas.Canvas(buffer)
    # A real text layer is more than one short line; spread the sentence over
    # consecutive lines so the scanned-page guard does not fire.
    words = text.split()
    for index, start in enumerate(range(0, len(words), 8)):
        chart.drawString(72, 760 - index * 16, " ".join(words[start:start + 8]))
    chart.showPage()
    chart.save()
    return buffer.getvalue()


def make_blank_pdf_bytes() -> bytes:
    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    chart = canvas.Canvas(buffer)
    chart.showPage()
    chart.save()
    return buffer.getvalue()


def make_encrypted_pdf_bytes() -> bytes:
    from pypdf import PdfReader, PdfWriter

    reader = PdfReader(io.BytesIO(make_pdf_bytes()))
    writer = PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
    writer.encrypt("secret")
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def make_docx_bytes() -> bytes:
    import docx

    document = docx.Document()
    document.add_heading("تنفس سلولی", level=1)
    document.add_paragraph("گلیکولیز در سیتوپلاسم رخ می‌دهد و دو ATP می‌سازد.")
    document.add_heading("چرخهٔ کربس", level=2)
    document.add_paragraph("چرخهٔ کربس در ماتریکس میتوکندری انجام می‌شود.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_image_only_docx_bytes() -> bytes:
    import docx

    document = docx.Document()
    document.add_paragraph("")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_pptx_bytes() -> bytes:
    from pptx import Presentation

    deck = Presentation()
    layout = deck.slide_layouts[1]
    first = deck.slides.add_slide(layout)
    first.shapes.title.text = "متابولیسم"
    first.placeholders[1].text = "گلوکز در گلیکولیز به پیروات تبدیل می‌شود."
    second = deck.slides.add_slide(layout)
    second.shapes.title.text = "زنجیرهٔ انتقال الکترون"
    second.placeholders[1].text = "بیشترین ATP در زنجیرهٔ انتقال الکترون ساخته می‌شود."
    deck.slides.add_slide(layout)  # an empty slide must not fail the extraction
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


class TempFileMixin(unittest.TestCase):
    def write(self, name: str, payload: bytes) -> Path:
        path = Path(self._tmp.name) / name
        path.write_bytes(payload)
        return path

    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)


class ClassifyTests(unittest.TestCase):
    def test_mime_types_win_and_extensions_are_the_fallback(self):
        self.assertEqual(classify_document("a.pdf", PDF_MIME), "pdf")
        self.assertEqual(classify_document("a.docx", DOCX_MIME), "docx")
        self.assertEqual(classify_document("a.pptx", PPTX_MIME), "pptx")
        # Some clients send application/octet-stream.
        self.assertEqual(
            classify_document("notes.pdf", "application/octet-stream"), "pdf"
        )
        self.assertEqual(classify_document("slides.PPT", None), "pptx")
        # Legacy .doc is accepted by extension; a friendly error follows later.
        self.assertEqual(classify_document("old.doc", None), "docx")
        self.assertIsNone(classify_document("photo.jpg", "image/jpeg"))
        self.assertIsNone(classify_document(None, None))

    def test_titles_are_clean_and_never_empty(self):
        self.assertEqual(document_title("  my  file.pdf ", "pdf"), "my file.pdf")
        self.assertEqual(document_title("", "pdf"), "سند PDF")
        self.assertEqual(document_title(None, "docx"), "سند Word")


class PdfExtractionTests(TempFileMixin):
    def test_text_layer_is_extracted_per_page_with_locators(self):
        path = self.write("notes.pdf", make_pdf_bytes())
        extracted = extract_document(path, "pdf")
        self.assertEqual(extracted.source_type, "pdf")
        self.assertEqual(extracted.page_count, 1)
        self.assertEqual(extracted.sections[0].locator, "صفحهٔ ۱")
        self.assertIn("ATP", extracted.text)

    def test_scanned_pdf_is_refused_with_the_ocr_message(self):
        path = self.write("scan.pdf", make_blank_pdf_bytes())
        with self.assertRaises(ScannedPdfError) as caught:
            extract_document(path, "pdf")
        self.assertIn("OCR پشتیبانی نمی‌شود", str(caught.exception))

    def test_password_protected_pdf_is_refused(self):
        path = self.write("locked.pdf", make_encrypted_pdf_bytes())
        with self.assertRaises(DocExtractError) as caught:
            extract_document(path, "pdf")
        self.assertIn("رمز", str(caught.exception))

    def test_threshold_is_a_floor_and_a_rate(self):
        self.assertTrue(looks_scanned([DocSection("صفحهٔ ۱", "", "کم")], 1))
        # One short line per slide/page is legitimate once the rate is met.
        sections = [DocSection(f"صفحهٔ {index}", "", "x" * 40) for index in range(10)]
        self.assertFalse(looks_scanned(sections, 10))


class DocxExtractionTests(TempFileMixin):
    def test_headings_become_locators(self):
        path = self.write("notes.docx", make_docx_bytes())
        extracted = extract_document(path, "docx")
        self.assertEqual(extracted.source_type, "docx")
        self.assertEqual(len(extracted.sections), 2)
        self.assertIn("تنفس سلولی", extracted.sections[0].locator)
        self.assertIn("گلیکولیز", extracted.text)
        self.assertIn("ماتریکس", extracted.text)

    def test_image_only_docx_is_refused(self):
        path = self.write("empty.docx", make_image_only_docx_bytes())
        with self.assertRaises(DocExtractError) as caught:
            extract_document(path, "docx")
        self.assertIn("OCR", str(caught.exception))


class PptxExtractionTests(TempFileMixin):
    def test_slides_are_extracted_with_slide_numbers(self):
        path = self.write("deck.pptx", make_pptx_bytes())
        extracted = extract_document(path, "pptx")
        self.assertEqual(extracted.source_type, "pptx")
        self.assertEqual(extracted.page_count, 3)
        locators = [section.locator for section in extracted.sections]
        self.assertIn("اسلاید ۱", locators)
        self.assertIn("اسلاید ۲", locators)
        # The empty third slide is skipped, not an error.
        self.assertNotIn("اسلاید ۳", locators)
        self.assertIn("گلیکولیز", extracted.text)

    def test_sections_from_slides_is_the_shared_shape(self):
        from types import SimpleNamespace

        slides = [
            SimpleNamespace(number=1, title="عنوان", body=("نکته",), notes="یادداشت"),
            SimpleNamespace(number=2, title="", body=(), notes=""),
        ]
        sections = sections_from_slides(slides)
        self.assertEqual(len(sections), 1)
        self.assertEqual(sections[0].locator, "اسلاید ۱")
        self.assertIn("یادداشت گوینده", sections[0].text)


class ChunkingTests(unittest.TestCase):
    def test_chunks_respect_the_budget_and_keep_locators(self):
        sections = [
            DocSection(f"صفحهٔ {index}", "", "م" * 400) for index in range(1, 6)
        ]
        chunks = chunk_sections(sections, max_chars=1000)
        self.assertGreater(len(chunks), 1)
        for chunk in chunks:
            self.assertLessEqual(len(chunk.text), 1000)
            self.assertTrue(chunk.locator)
            self.assertTrue(chunk.sources)

    def test_a_chunk_carries_the_pages_it_spans(self):
        sections = [
            DocSection("صفحهٔ ۱", "", "الف" * 100),
            DocSection("صفحهٔ ۲", "", "ب" * 100),
        ]
        chunks = chunk_sections(sections, max_chars=1000)
        self.assertEqual(len(chunks), 1)
        self.assertEqual(chunks[0].sources, ("صفحهٔ ۱", "صفحهٔ ۲"))
        self.assertIn("تا", chunks[0].locator)

    def test_an_oversized_section_is_split_on_boundaries(self):
        paragraph = "جملهٔ بلند. " * 200
        chunks = chunk_sections([DocSection("صفحهٔ ۷", "", paragraph)], max_chars=600)
        self.assertGreater(len(chunks), 1)
        for index, chunk in enumerate(chunks, start=1):
            self.assertLessEqual(len(chunk.text), 700)
            # Locators are user-facing, so the fraction is in Persian digits.
            self.assertIn(
                f"({to_persian_digits(index)}/{to_persian_digits(len(chunks))})",
                chunk.locator,
            )

    def test_flat_text_is_labelled_by_part(self):
        chunks = chunk_text("متن " * 2000, max_chars=MAX_MODEL_INPUT_CHARS)
        self.assertGreaterEqual(len(chunks), 1)
        self.assertTrue(all("بخش" in chunk.locator for chunk in chunks))

    def test_json_round_trip_and_repair(self):
        chunks = chunk_text("یک متن کوتاه برای آزمون", max_chars=MAX_MODEL_INPUT_CHARS)
        restored = chunks_from_json(chunks_to_json(chunks))
        self.assertEqual([chunk.text for chunk in restored], [c.text for c in chunks])
        self.assertEqual(restored[0].locator, chunks[0].locator)
        # A corrupt cache is ignored instead of raising.
        self.assertEqual(chunks_from_json("{not json"), [])
        self.assertEqual(chunks_from_json(None), [])


class HelperTests(TempFileMixin):
    def test_sha256_is_stable_and_distinguishes_files(self):
        first = self.write("a.bin", b"alpha")
        second = self.write("b.bin", b"beta")
        self.assertEqual(sha256_file(first), sha256_file(first))
        self.assertNotEqual(sha256_file(first), sha256_file(second))

    def test_preview_is_word_safe(self):
        self.assertEqual(preview_text("کوتاه"), "کوتاه")
        long = "کلمه " * 200
        preview = preview_text(long, limit=30)
        self.assertLessEqual(len(preview), 31)
        self.assertTrue(preview.endswith("…"))

    def test_doc_chunk_is_hashable_data(self):
        chunk = DocChunk(index=0, locator="صفحهٔ ۱", sources=("صفحهٔ ۱",), text="متن")
        self.assertEqual(chunk.index, 0)
        self.assertIn("صفحهٔ ۱", chunk.sources)


if __name__ == "__main__":
    unittest.main()
