"""Extract narration audio and slide text from PowerPoint decks.

A .pptx file is an OPC (ZIP) package: narration lives in ``ppt/media/*`` and is
bound to a slide through that slide's ``_rels`` part. The reader below walks the
package with the standard library so it stays in control of every byte it
extracts (path traversal, entry count and unpacked size are all bounded), and
uses python-pptx only for the slide/notes text.
"""

from __future__ import annotations

import asyncio
import logging
import math
import posixpath
import re
import zipfile
from bisect import bisect_right
from dataclasses import dataclass, replace
from pathlib import Path, PurePosixPath

from lxml import etree

from .config import Settings
from .media import merge_audio_tracks, needs_transcode, probe_media_many
from .stt import largest_configured_upload
from .transcript import RawTranscript, Word, build_document

logger = logging.getLogger(__name__)

NATIVE_EXTENSIONS = {".pptx", ".pptm", ".ppsx", ".ppsm", ".potx", ".potm"}
LEGACY_EXTENSIONS = {".ppt", ".pps", ".pot"}
# ODP/OTP have no verified pure-Python converter with acceptable fidelity, so
# they are classified as explicitly unsupported instead of pretending they work.
UNSUPPORTED_EXTENSIONS = {".odp", ".otp"}
PRESENTATION_MIME_TYPES = {
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": "native",
    "application/vnd.openxmlformats-officedocument.presentationml.slideshow": "native",
    "application/vnd.openxmlformats-officedocument.presentationml.template": "native",
    "application/vnd.ms-powerpoint.presentation.macroenabled.12": "native",
    "application/vnd.ms-powerpoint.slideshow.macroenabled.12": "native",
    "application/vnd.ms-powerpoint.template.macroenabled.12": "native",
    "application/vnd.oasis.opendocument.presentation-template": "unsupported",
    "application/vnd.ms-powerpoint": "legacy",
    "application/mspowerpoint": "legacy",
    "application/powerpoint": "legacy",
    "application/vnd.oasis.opendocument.presentation": "unsupported",
}

AUDIO_MEDIA_EXTENSIONS = {
    ".aac", ".aiff", ".aif", ".amr", ".au", ".flac", ".m4a", ".m4b", ".mp3",
    ".oga", ".ogg", ".opus", ".wav", ".wma",
}
VIDEO_MEDIA_EXTENSIONS = {
    ".3gp", ".asf", ".avi", ".m4v", ".mkv", ".mov", ".mp4", ".mpeg", ".mpg",
    ".webm", ".wmv",
}

RELS_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
PML_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
DRAWING_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
OFFICE_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
MS_OFFICE_REL_NS = "http://schemas.microsoft.com/office/2007/relationships"
MEDIA_REL_TYPES = {
    f"{OFFICE_REL_NS}/audio": "audio",
    f"{OFFICE_REL_NS}/video": "video",
    f"{OFFICE_REL_NS}/media": "media",
    f"{MS_OFFICE_REL_NS}/media": "media",
}
NOTES_SLIDE_REL_TYPE = f"{OFFICE_REL_NS}/notesSlide"
#: Maximum STT timestamp jitter (seconds) tolerated at a clip boundary when a
#: word overlaps only that clip and at least half of the word lies inside it.
ALIGNMENT_BOUNDARY_TOLERANCE_SECONDS = 0.08
#: Placeholders that carry page furniture (slide number, date, header, footer),
#: never the lecture's own words; python-pptx skips their ``a:fld`` text too.
FURNITURE_PLACEHOLDERS = frozenset({"sldNum", "hdr", "ftr", "dt"})
#: ``type`` values of the placeholders that hold a slide's title.
TITLE_PLACEHOLDERS = frozenset({"title", "ctrTitle"})
DIAGRAM_NS = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
CHART_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"
#: SmartArt keeps its words in a diagram-data part, a chart in a chart part;
#: neither is a text frame, so python-pptx (and a naive slide-XML scan) misses
#: them. Both parts are read and merged back into their slide's text.
DIAGRAM_DATA_REL_TYPE = f"{OFFICE_REL_NS}/diagramData"
CHART_REL_TYPE = f"{OFFICE_REL_NS}/chart"
MAX_ZIP_ENTRIES = 20000
MAX_ZIP_MEMBER_NAME_BYTES = 4096
MAX_ZIP_CENTRAL_DIRECTORY_BYTES = 32 * 1024 * 1024


class PresentationError(RuntimeError):
    """Raised for decks the bot cannot or should not process."""


@dataclass(frozen=True, slots=True)
class SlideText:
    number: int
    title: str = ""
    body: tuple[str, ...] = ()
    notes: str = ""

    @property
    def is_empty(self) -> bool:
        return not (self.title.strip() or any(self.body) or self.notes.strip())

    def as_markdown(self, include_notes: bool = True) -> str:
        heading = f"### اسلاید {self.number}" + (f" — {self.title.strip()}" if self.title.strip() else "")
        lines = [heading]
        for raw_item in self.body:
            item = raw_item.strip()
            if not item:
                continue
            if " | " in item:
                row = item if item.startswith("|") else f"| {item} |"
                lines.append(row)
            else:
                lines.append(f"- {item}")
        if include_notes and self.notes.strip():
            note_paragraphs = [
                " ".join(part.split())
                for part in self.notes.splitlines()
                if part.strip()
            ]
            if note_paragraphs:
                if len(lines) > 1:
                    lines.append("")
                lines.append("یادداشت گوینده: " + note_paragraphs[0])
                lines.extend(note_paragraphs[1:])
        return "\n".join(lines)


@dataclass(frozen=True, slots=True)
class MediaClip:
    order: int
    slide_number: int | None
    part_name: str
    path: Path
    kind: str
    duration: float | None = None
    relationship_id: str | None = None

    @property
    def label(self) -> str:
        where = f"اسلاید {self.slide_number}" if self.slide_number else "بدون اسلاید"
        return f"{where} ({posixpath.basename(self.part_name)})"


@dataclass(frozen=True, slots=True)
class ClipSpan:
    """The interval a source clip occupies in the merged STT audio."""

    clip: MediaClip
    start_seconds: float
    end_seconds: float

    @property
    def duration(self) -> float:
        return max(0.0, self.end_seconds - self.start_seconds)


@dataclass(frozen=True, slots=True)
class PresentationContent:
    slides: tuple[SlideText, ...]
    clips: tuple[MediaClip, ...]
    skipped: tuple[str, ...] = ()
    total_slides: int | None = None

    @property
    def slide_count(self) -> int:
        return self.total_slides if self.total_slides is not None else len(self.slides)


@dataclass(frozen=True, slots=True)
class PreparedAudio:
    path: Path
    clips: tuple[MediaClip, ...]
    source_duration: float
    output_duration: float
    clip_spans: tuple[ClipSpan, ...]
    merged: bool
    skipped: tuple[str, ...] = ()


@dataclass(frozen=True, slots=True)
class PresentationNarration:
    """Timed narration mapped only where a clip-to-slide relation is known."""

    slide_text: tuple[tuple[int, str], ...]
    unassigned_text: str
    aligned: bool

    def as_text(self) -> str:
        if not self.aligned:
            text = self.unassigned_text.strip()
            if not text:
                return ""
            return "### گفتار بدون نگاشت مطمئن به اسلاید\n\n" + text
        blocks = [
            f"### گفتار مرتبط با اسلاید {number}\n\n{text}"
            for number, text in self.slide_text
            if text.strip()
        ]
        if self.unassigned_text.strip():
            blocks.append(
                "### گفتار بدون نگاشت مطمئن به اسلاید\n\n" + self.unassigned_text.strip()
            )
        return "\n\n".join(blocks)


def classify_presentation(filename: str | None, mime_type: str | None) -> str | None:
    """Return 'native', 'legacy', 'unsupported' or None for a Telegram document."""
    suffix = Path(filename).suffix.lower() if filename else ""
    if suffix in NATIVE_EXTENSIONS:
        return "native"
    if suffix in LEGACY_EXTENSIONS:
        return "legacy"
    if suffix in UNSUPPORTED_EXTENSIONS:
        return "unsupported"
    normalized = (mime_type or "").split(";")[0].strip().lower()
    return PRESENTATION_MIME_TYPES.get(normalized)


def media_kind(part_name: str) -> str | None:
    suffix = posixpath.splitext(part_name)[1].lower()
    if suffix in AUDIO_MEDIA_EXTENSIONS:
        return "audio"
    if suffix in VIDEO_MEDIA_EXTENSIONS:
        return "video"
    return None


def _resolve_target(base_part: str, target: str) -> str:
    base_dir = posixpath.dirname(base_part)
    joined = target if target.startswith("/") else posixpath.join(base_dir, target)
    return posixpath.normpath(joined).lstrip("/")


def _rels_part_for(part_name: str) -> str:
    directory, name = posixpath.split(part_name)
    return posixpath.join(directory, "_rels", f"{name}.rels")


_XML_PARSER = etree.XMLParser(
    resolve_entities=False, no_network=True, huge_tree=False, recover=False
)


def _parse_xml(payload: bytes) -> etree._Element:
    return etree.fromstring(payload, parser=_XML_PARSER)


def _zip_member_stream_complete(handle, info: zipfile.ZipInfo) -> bool:
    """Check that a member's compressed stream was consumed to its true EOF.

    ``ZipExtFile`` stops at the central-directory ``file_size``. Checking its
    stream state catches a forged small ``file_size`` that would otherwise
    make a ZIP bomb appear harmless while leaving compressed payload unread.
    """
    if getattr(handle, "_compress_left", None) != 0:
        return False
    if info.compress_type == zipfile.ZIP_STORED:
        return info.compress_size == info.file_size
    decompressor = getattr(handle, "_decompressor", None)
    return bool(
        decompressor is not None
        and decompressor.eof
        and not decompressor.unconsumed_tail
        and not decompressor.unused_data
    )


def _read_member(archive: zipfile.ZipFile, name: str, limit: int) -> bytes | None:
    """Read one ZIP member under a hard output limit; never silently truncate."""
    try:
        info = archive.getinfo(name)
    except KeyError:
        return None
    if limit < 0 or info.file_size < 0 or info.file_size > limit:
        raise PresentationError("یکی از بخش‌های فایل ارائه بیش از حد بزرگ است.")
    try:
        with archive.open(info) as handle:
            # ZipExtFile normally stops at central-directory file_size. Permit
            # one byte beyond that declaration (but never beyond our limit) so
            # a forged smaller value is detected instead of silently hiding
            # decompressed bytes.
            member_cap = min(limit, info.file_size)
            handle._left = member_cap + 1
            payload = handle.read(member_cap + 1)
            complete = _zip_member_stream_complete(handle, info)
    except (zipfile.BadZipFile, EOFError, OSError, RuntimeError) as exc:
        raise PresentationError("یکی از بخش‌های فایل ارائه خراب است.") from exc
    if len(payload) > limit or len(payload) != info.file_size or not complete:
        raise PresentationError("اندازهٔ واقعی یکی از بخش‌های فایل ارائه معتبر نیست.")
    return payload


def _relationships(archive: zipfile.ZipFile, part_name: str) -> list[tuple[str, str, str]]:
    """Return (id, type, internal_target) triples; external targets are blank."""
    payload = _read_member(archive, _rels_part_for(part_name), 8_000_000)
    if not payload:
        return []
    try:
        root = _parse_xml(payload)
    except etree.XMLSyntaxError:
        logger.warning("Malformed relationship part: %s", _safe_log_name(part_name))
        return []
    result = []
    for element in root.findall(f"{{{RELS_NS}}}Relationship"):
        rel_id = element.get("Id") or ""
        rel_type = element.get("Type") or ""
        target = element.get("Target") or ""
        mode = element.get("TargetMode") or "Internal"
        if rel_id and target:
            result.append((rel_id, rel_type, target if mode == "Internal" else ""))
    return result


def _slide_parts_in_order(archive: zipfile.ZipFile) -> list[str]:
    """Slide part names following the deck's own slide order."""
    payload = _read_member(archive, "ppt/presentation.xml", 20_000_000)
    # Built once: ``namelist()`` rebuilds a list on every call, which turns the
    # membership test below into a quadratic scan on large decks.
    names = set(archive.namelist())
    rels = {
        rel_id: _resolve_target("ppt/presentation.xml", target)
        for rel_id, _type, target in _relationships(archive, "ppt/presentation.xml")
        if target
    }
    ordered: list[str] = []
    if payload:
        try:
            root = _parse_xml(payload)
        except etree.XMLSyntaxError as exc:
            raise PresentationError("ساختار XML فایل ارائه معتبر نیست.") from exc
        for node in root.iterfind(f".//{{{PML_NS}}}sldIdLst/{{{PML_NS}}}sldId"):
            target = rels.get(node.get(f"{{{OFFICE_REL_NS}}}id") or "")
            if target and target in names and target not in ordered:
                ordered.append(target)
    if ordered:
        return ordered
    # Fall back to numeric filename ordering when the deck omits sldIdLst.
    def slide_index(name: str) -> tuple[int, str]:
        match = re.search(r"slide(\d+)\.xml$", name)
        return (int(match.group(1)) if match else 1 << 30, name)

    return sorted(
        (n for n in names if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)),
        key=slide_index,
    )


def natural_key(name: str) -> tuple:
    """Sort media2.m4a before media10.m4a instead of lexicographically."""
    # The empty pieces are kept on purpose: re.split alternates text/number, so
    # every tuple position keeps one type and comparisons stay well defined.
    return tuple(
        int(part) if part.isdigit() else part.lower()
        for part in re.split(r"(\d+)", name)
    )


def _safe_log_name(name: object, *, limit: int = 200) -> str:
    """A deck-controlled part name that is safe to put on one log line."""
    text = " ".join(str(name).split())
    return text[:limit] + ("…" if len(text) > limit else "")


def _safe_media_filename(order: int, part_name: str) -> str:
    base = posixpath.basename(part_name) or "media"
    cleaned = re.sub(r"[^A-Za-z0-9._-]", "_", base)[-64:].lstrip(".") or "media"
    return f"{order:04d}-{cleaned}"


def _is_zip_symlink(info: zipfile.ZipInfo) -> bool:
    # POSIX symlink is stored with file type 0xA000 in external_attr
    return (info.external_attr >> 16) & 0xA000 == 0xA000


def _validate_archive(archive: zipfile.ZipFile, settings: Settings) -> None:
    infos = archive.infolist()
    if len(infos) > MAX_ZIP_ENTRIES:
        raise PresentationError("تعداد بخش‌های داخل فایل ارائه بیش از حد مجاز است.")
    total = 0
    for info in infos:
        name = info.filename.replace("\\", "/")
        parts = PurePosixPath(name).parts
        if name.startswith("/") or ".." in parts or PurePosixPath(name).is_absolute():
            raise PresentationError("فایل ارائه شامل مسیر نامعتبر است و پردازش نشد.")
        if len(name.encode("utf-8", "surrogatepass")) > MAX_ZIP_MEMBER_NAME_BYTES:
            raise PresentationError("نام یکی از بخش‌های فایل ارائه بیش از حد طولانی است.")
        if _is_zip_symlink(info):
            raise PresentationError("فایل ارائه شامل پیوند نمادین است و پردازش نشد.")
        # Control characters would let a crafted part name forge extra lines
        # in the line-oriented log (audit F-19).
        if any(ord(ch) < 0x20 or ord(ch) == 0x7F for ch in name):
            raise PresentationError("فایل ارائه شامل نام بخش نامعتبر است و پردازش نشد.")
        if info.file_size < 0 or info.compress_size < 0:
            raise PresentationError("اندازهٔ یکی از بخش‌های فایل ارائه معتبر نیست.")
        if info.compress_type not in {zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED}:
            raise PresentationError("روش فشرده‌سازی یکی از بخش‌های ارائه پشتیبانی نمی‌شود.")
        total += info.file_size
        if total > settings.presentation_max_unpacked_bytes:
            raise PresentationError("حجم بازشدهٔ فایل ارائه از حد مجاز بیشتر است.")


#: PowerPoint stores a soft line break inside one run as U+000B and python-pptx
#: passes it through in ``shape.text``. It is not meaningful to the note model
#: (and some providers reject raw control characters), so it is folded into a
#: space while the text is still on its way in.
_SOFT_BREAKS = str.maketrans({"\v": " ", "\u2028": " ", "\u2029": " "})


def _clean_text(text: str) -> str:
    """One line of clean prose: no control characters, no stray whitespace."""
    return " ".join(text.translate(_SOFT_BREAKS).split())


def _extract_slide_text_with_pptx(path: Path) -> list[SlideText]:
    """Read titles, bullet text, tables and speaker notes with python-pptx."""
    try:
        # Presentation() rejects slideshow/template main content types even
        # though their slide structure is identical. Register the two missing
        # macro-enabled variants and open the package directly (no macros run).
        from pptx.opc.package import PartFactory
        from pptx.package import Package
        from pptx.parts.presentation import PresentationPart
    except ImportError:  # pragma: no cover - dependency is declared in requirements
        logger.warning("python-pptx is not installed; slide text will be skipped")
        return []
    try:
        main_types = {
            "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml",
            "application/vnd.openxmlformats-officedocument.presentationml.slideshow.main+xml",
            "application/vnd.openxmlformats-officedocument.presentationml.template.main+xml",
            "application/vnd.ms-powerpoint.presentation.macroEnabled.main+xml",
            "application/vnd.ms-powerpoint.slideshow.macroEnabled.main+xml",
            "application/vnd.ms-powerpoint.template.macroEnabled.main+xml",
        }
        for content_type in main_types:
            PartFactory.part_type_for.setdefault(content_type, PresentationPart)
        main_part = Package.open(str(path)).main_document_part
        if main_part.content_type not in main_types:
            raise PresentationError("نوع فایل ارائه پشتیبانی نمی‌شود.")
        deck = main_part.presentation
    except Exception:
        logger.exception("python-pptx could not open the deck; continuing without slide text")
        return []

    def shape_texts(shapes) -> list[str]:
        collected: list[str] = []
        for shape in shapes:
            try:
                if getattr(shape, "shape_type", None) is not None and shape.shape_type == 6:
                    collected.extend(shape_texts(shape.shapes))  # grouped shapes
                    continue
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        cells = [_clean_text(cell.text) for cell in row.cells]
                        if any(cells):
                            line = " | ".join(cells) if len(cells) > 1 else cells[0]
                            if line:
                                collected.append(line)
                    continue
                if getattr(shape, "has_text_frame", False):
                    for paragraph in shape.text_frame.paragraphs:
                        line = _clean_text(
                            "".join(run.text for run in paragraph.runs)
                        )
                        if line:
                            collected.append(line)
            except Exception:  # a single broken shape must not lose the deck
                logger.debug("Unreadable shape skipped", exc_info=True)
        return collected

    slides: list[SlideText] = []
    for number, slide in enumerate(deck.slides, start=1):
        title = ""
        title_lines: set[str] = set()
        try:
            if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
                title_shape = slide.shapes.title
                title_parts = [
                    _clean_text("".join(run.text for run in paragraph.runs))
                    for paragraph in title_shape.text_frame.paragraphs
                ]
                title_parts = [part for part in title_parts if part]
                title_lines.update(title_parts)
                title = " ".join(title_parts) or _clean_text(title_shape.text)
                if title:
                    title_lines.add(title)
        except Exception:
            logger.debug("Slide %s has no readable title", number, exc_info=True)
        body = [line for line in shape_texts(slide.shapes) if line and line not in title_lines]
        notes = ""
        try:
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                note_lines = [
                    _clean_text("".join(run.text for run in paragraph.runs))
                    for paragraph in slide.notes_slide.notes_text_frame.paragraphs
                ]
                note_lines = [line for line in dict.fromkeys(note_lines) if line]
                notes = "\n".join(note_lines) or _clean_text(slide.notes_slide.notes_text_frame.text)
        except Exception:
            logger.debug("Slide %s has no readable notes", number, exc_info=True)
        slides.append(SlideText(number, title, tuple(dict.fromkeys(body)), notes))
    return slides


def _local_name(element) -> str:
    """The local part of a namespaced tag (``{ns}sp`` → ``sp``)."""
    tag = element.tag
    if not isinstance(tag, str):
        return ""
    return tag.rsplit("}", 1)[-1]


def _paragraph_text(paragraph) -> str:
    """The visible text of one ``a:p``, in document order.

    Only ``a:r`` runs and ``a:br`` breaks are read. An ``a:fld`` (a slide
    number, a date) is page furniture and python-pptx's ``paragraph.runs``
    skips it too, so skipping it here keeps the two readers in step.
    """
    parts: list[str] = []
    for node in paragraph:
        if node.tag == f"{{{DRAWING_NS}}}r":
            parts.extend(t.text or "" for t in node.iter(f"{{{DRAWING_NS}}}t"))
        elif node.tag == f"{{{DRAWING_NS}}}br":
            # python-pptx renders ``a:br`` as U+000B; keep the word boundary.
            parts.append(" ")
    return _clean_text("".join(parts))


def _text_body_lines(shape) -> list[str]:
    """The non-empty paragraphs of a shape's ``p:txBody``."""
    body = shape.find(f"{{{PML_NS}}}txBody")
    if body is None:
        return []
    lines = [
        _paragraph_text(paragraph)
        for paragraph in body.findall(f"{{{DRAWING_NS}}}p")
    ]
    return [line for line in lines if line]


def _placeholder_type(shape) -> str:
    """The ``p:ph/@type`` of a placeholder shape, or ``""`` for a plain shape."""
    node = shape.find(f"{{{PML_NS}}}nvSpPr/{{{PML_NS}}}nvPr/{{{PML_NS}}}ph")
    return (node.get("type") or "") if node is not None else ""


def _table_lines(frame) -> list[str]:
    """One line per table row, cells joined the way the python-pptx reader does."""
    lines: list[str] = []
    for table in frame.iter(f"{{{DRAWING_NS}}}tbl"):
        for row in table.findall(f"{{{DRAWING_NS}}}tr"):
            cells = [
                " ".join(
                    line
                    for line in (
                        _paragraph_text(paragraph)
                        for paragraph in cell.iter(f"{{{DRAWING_NS}}}p")
                    )
                    if line
                )
                for cell in row.findall(f"{{{DRAWING_NS}}}tc")
            ]
            if any(cells):
                line = " | ".join(cells) if len(cells) > 1 else cells[0]
                if line:
                    lines.append(line)
    return lines


def _shape_tree_lines(node) -> list[tuple[bool, str]]:
    """``(is_title, line)`` pairs of one ``p:spTree``/``p:grpSp``, in order."""
    entries: list[tuple[bool, str]] = []
    for child in node:
        name = _local_name(child)
        if name == "sp":
            is_title = _placeholder_type(child) in TITLE_PLACEHOLDERS
            entries.extend((is_title, line) for line in _text_body_lines(child))
        elif name == "grpSp":
            entries.extend(_shape_tree_lines(child))
        elif name == "AlternateContent":
            # ``mc:AlternateContent`` wraps a newer shape plus a fallback of
            # the same shape; descend into both, the duplicate is removed by
            # the caller's de-duplication.
            for branch in child:
                entries.extend(_shape_tree_lines(branch))
        elif name == "graphicFrame":
            entries.extend((False, line) for line in _table_lines(child))
    return entries


def _related_part(
    archive: zipfile.ZipFile,
    source_part: str,
    rel_id: str,
    rel_type: str,
    *,
    rels: list[tuple[str, str, str]] | None = None,
) -> str | None:
    """The internal part ``rel_id`` of ``source_part`` names, if it matches."""
    if not rel_id:
        return None
    entries = rels if rels is not None else _relationships(archive, source_part)
    for candidate_id, candidate_type, target in entries:
        if candidate_id == rel_id and candidate_type == rel_type and target:
            name = _resolve_target(source_part, target)
            try:
                archive.getinfo(name)
            except KeyError:
                return None
            return name
    return None


def _lines_of(element) -> list[str]:
    """One line per paragraph of an element's drawing text."""
    return [
        line
        for line in (_paragraph_text(p) for p in element.iter(f"{{{DRAWING_NS}}}p"))
        if line
    ]


def _diagram_lines(archive: zipfile.ZipFile, part_name: str) -> list[str]:
    """The words of a SmartArt diagram, from its ``dgm:dataModel`` part.

    SmartArt text is not a shape text frame: it lives in ``ppt/diagrams`` and
    python-pptx never sees it, so a diagram-only slide used to read as empty.
    """
    payload = _read_member(archive, part_name, 8_000_000)
    if not payload:
        return []
    try:
        root = _parse_xml(payload)
    except etree.XMLSyntaxError:
        logger.warning("Malformed diagram part: %s", _safe_log_name(part_name))
        return []
    lines: list[str] = []
    for point in root.iter(f"{{{DIAGRAM_NS}}}pt"):
        if point.get("type"):
            continue  # doc/transition nodes are wiring, not content
        text = point.find(f"{{{DIAGRAM_NS}}}t")
        if text is None:
            continue
        lines.extend(_lines_of(text))
        if len(lines) >= 200:
            lines.append("…")
            break
    return lines


def _chart_lines(archive: zipfile.ZipFile, part_name: str) -> list[str]:
    """Read a chart's title, labels, series names and bounded cached values."""
    payload = _read_member(archive, part_name, 8_000_000)
    if not payload:
        return []
    try:
        root = _parse_xml(payload)
    except etree.XMLSyntaxError:
        logger.warning("Malformed chart part: %s", _safe_log_name(part_name))
        return []
    lines: list[str] = []
    for chart in root.iter(f"{{{CHART_NS}}}chart"):
        titles: list[str] = []
        for title in chart.findall(f"{{{CHART_NS}}}title"):
            titles.extend(_lines_of(title))
        if titles:
            lines.append("عنوان نمودار: " + " ".join(dict.fromkeys(titles)))

        category_by_index: dict[int, str] = {}
        category_order: list[str] = []
        for category in chart.iter(f"{{{CHART_NS}}}cat"):
            for position, point in enumerate(category.iter(f"{{{CHART_NS}}}pt")):
                value = point.find(f"{{{CHART_NS}}}v")
                text = (value.text or "").strip() if value is not None else ""
                if not text:
                    continue
                try:
                    index = int(point.get("idx", position))
                except (TypeError, ValueError):
                    index = position
                category_by_index[index] = text
                category_order.append(text)
        category_order = list(dict.fromkeys(category_order))
        if category_order:
            shown = category_order[:60]
            if len(category_order) > 60:
                shown.append("…")
            lines.append("دسته‌های نمودار: " + "، ".join(shown))

        series_names: list[str] = []
        for entry in chart.iter(f"{{{CHART_NS}}}ser"):
            name_node = entry.find(f"{{{CHART_NS}}}tx")
            name_values: list[str] = []
            if name_node is not None:
                name_values.extend(_lines_of(name_node))
                name_values.extend(
                    (value.text or "").strip()
                    for value in name_node.iter(f"{{{CHART_NS}}}v")
                    if (value.text or "").strip()
                )
            name = " ".join(dict.fromkeys(name_values)).strip()
            if name:
                series_names.append(name)

            value_node = entry.find(f"{{{CHART_NS}}}val")
            if value_node is None:
                continue
            points: list[tuple[int, str]] = []
            for position, point in enumerate(value_node.iter(f"{{{CHART_NS}}}pt")):
                value = point.find(f"{{{CHART_NS}}}v")
                text = (value.text or "").strip() if value is not None else ""
                if not text:
                    continue
                try:
                    index = int(point.get("idx", position))
                except (TypeError, ValueError):
                    index = position
                points.append((index, text))
            if points:
                shown_points = points[:60]
                value_labels = [
                    f"{category_by_index[index]}={value}"
                    if index in category_by_index else value
                    for index, value in shown_points
                ]
                if len(points) > 60:
                    value_labels.append("…")
                label = f"مقادیر سری «{name}»: " if name else "مقادیر نمودار: "
                lines.append(label + "، ".join(value_labels))
        if series_names:
            lines.append("سری‌های نمودار: " + "، ".join(dict.fromkeys(series_names)))
        break  # a chart part holds exactly one c:chart
    return lines


def _graphic_frame_lines(
    archive: zipfile.ZipFile,
    slide_part: str,
    node,
    *,
    rels: list[tuple[str, str, str]] | None = None,
) -> list[str]:
    """Diagram/chart text of every ``p:graphicFrame`` under ``node``.

    A SmartArt's words live in its diagram-data part and a chart's labels in
    its chart part; neither is a shape text frame, so both are resolved
    through the slide's relationships and read from their own parts.
    """
    lines: list[str] = []
    for frame in node.iter(f"{{{PML_NS}}}graphicFrame"):
        for data in frame.iter(f"{{{DRAWING_NS}}}graphicData"):
            for child in data:
                if child.tag == f"{{{DIAGRAM_NS}}}relIds":
                    part = _related_part(
                        archive,
                        slide_part,
                        child.get(f"{{{OFFICE_REL_NS}}}dm") or "",
                        DIAGRAM_DATA_REL_TYPE,
                        rels=rels,
                    )
                    if part:
                        lines.extend(_diagram_lines(archive, part))
                elif child.tag == f"{{{CHART_NS}}}chart":
                    part = _related_part(
                        archive,
                        slide_part,
                        child.get(f"{{{OFFICE_REL_NS}}}id") or "",
                        CHART_REL_TYPE,
                        rels=rels,
                    )
                    if part:
                        lines.extend(_chart_lines(archive, part))
    return list(dict.fromkeys(lines))


def _notes_part_for(
    archive: zipfile.ZipFile,
    slide_part: str,
    *,
    rels: list[tuple[str, str, str]] | None = None,
) -> str | None:
    """The notes part a slide points at through its relationships, if any."""
    entries = rels if rels is not None else _relationships(archive, slide_part)
    for _rel_id, rel_type, target in entries:
        if rel_type == NOTES_SLIDE_REL_TYPE and target:
            return _resolve_target(slide_part, target)
    return None


def _notes_text_from_xml(
    archive: zipfile.ZipFile,
    slide_part: str,
    *,
    rels: list[tuple[str, str, str]] | None = None,
) -> str:
    """Speaker-note text of one slide, read from ``ppt/notesSlides/*.xml``."""
    notes_part = _notes_part_for(archive, slide_part, rels=rels)
    if not notes_part:
        return ""
    payload = _read_member(archive, notes_part, 8_000_000)
    if not payload:
        return ""
    try:
        root = _parse_xml(payload)
    except etree.XMLSyntaxError:
        logger.warning("Malformed notes part: %s", _safe_log_name(notes_part))
        return ""
    tree = root.find(f"{{{PML_NS}}}cSld/{{{PML_NS}}}spTree")
    if tree is None:
        return ""
    lines: list[str] = []
    for shape in tree.iter(f"{{{PML_NS}}}sp"):
        if _placeholder_type(shape) in FURNITURE_PLACEHOLDERS:
            continue
        lines.extend(_text_body_lines(shape))
    return "\n".join(dict.fromkeys(lines))


def _slide_text_from_xml(
    archive: zipfile.ZipFile,
    slide_part: str,
    number: int,
    *,
    rels: list[tuple[str, str, str]] | None = None,
) -> SlideText:
    """One slide's title, body, tables and speaker notes from its XML part."""
    payload = _read_member(archive, slide_part, 20_000_000)
    if not payload:
        return SlideText(number)
    try:
        root = _parse_xml(payload)
    except etree.XMLSyntaxError:
        logger.warning("Malformed slide part: %s", _safe_log_name(slide_part))
        return SlideText(number)
    slide_rels = rels if rels is not None else _relationships(archive, slide_part)
    tree = root.find(f"{{{PML_NS}}}cSld/{{{PML_NS}}}spTree")
    entries = _shape_tree_lines(tree) if tree is not None else []
    title = " ".join(line for is_title, line in entries if is_title).strip()
    body = [
        line
        for is_title, line in entries
        if not is_title and line and line != title
    ]
    body.extend(
        line
        for line in _graphic_frame_lines(archive, slide_part, root, rels=slide_rels)
        if line and line != title
    )
    return SlideText(
        number,
        title,
        tuple(dict.fromkeys(body)),
        _notes_text_from_xml(archive, slide_part, rels=slide_rels),
    )


def _extract_slide_text_from_xml(
    path: Path,
    *,
    archive: zipfile.ZipFile | None = None,
    slide_parts: list[str] | None = None,
    slide_rels_by_part: dict[str, list[tuple[str, str, str]]] | None = None,
) -> list[SlideText]:
    """Read every slide's text straight from the OPC package (no python-pptx).

    This is the fallback for a deck python-pptx cannot open (or an environment
    where it is not importable): the words on the slides still reach the note
    model instead of the narration being the only input. Only text stored as
    ``a:t`` is read — text that exists only as a picture stays a picture.
    """
    if archive is not None:
        try:
            parts = slide_parts if slide_parts is not None else _slide_parts_in_order(archive)
            return [
                _slide_text_from_xml(
                    archive,
                    part,
                    number,
                    rels=(slide_rels_by_part or {}).get(part),
                )
                for number, part in enumerate(parts, start=1)
            ]
        except Exception:  # a damaged package must not cost the narration
            logger.warning("could not read slide text from the package XML", exc_info=True)
            return []
    if not zipfile.is_zipfile(path):
        return []
    try:
        with zipfile.ZipFile(path) as opened_archive:
            parts = slide_parts if slide_parts is not None else _slide_parts_in_order(opened_archive)
            return [
                _slide_text_from_xml(
                    opened_archive,
                    part,
                    number,
                    rels=(slide_rels_by_part or {}).get(part),
                )
                for number, part in enumerate(parts, start=1)
            ]
    except Exception:  # a damaged package must not cost the narration
        logger.warning("could not read slide text from the package XML", exc_info=True)
        return []


def _extract_slide_text(
    path: Path,
    *,
    archive: zipfile.ZipFile | None = None,
    slide_parts: list[str] | None = None,
    slide_rels_by_part: dict[str, list[tuple[str, str, str]]] | None = None,
) -> list[SlideText]:
    """Merge python-pptx text with XML fallback on each individual slide.

    The XML reader runs even when python-pptx succeeds: a deck can be only
    partially readable (one broken slide, a table/notes edge case, or text held
    in SmartArt/chart parts). We keep the richer primary result and fill every
    missing field from that slide's package XML instead of using an all-or-none
    fallback for the entire deck.
    """
    primary = _extract_slide_text_with_pptx(path)
    fallback = _extract_slide_text_from_xml(
        path,
        archive=archive,
        slide_parts=slide_parts,
        slide_rels_by_part=slide_rels_by_part,
    )
    primary_by_number = {slide.number: slide for slide in primary}
    fallback_by_number = {slide.number: slide for slide in fallback}
    numbers = sorted(set(primary_by_number) | set(fallback_by_number))
    merged: list[SlideText] = []
    recovered = 0

    for number in numbers:
        preferred = primary_by_number.get(number, SlideText(number))
        xml = fallback_by_number.get(number, SlideText(number))
        if preferred.is_empty and not xml.is_empty:
            merged.append(xml)
            recovered += 1
            continue

        title = preferred.title or xml.title
        body = list(preferred.body)
        seen = {line for line in (title, *body) if line}
        # If both extractors found different candidate titles, preserve the
        # secondary text as body rather than silently discarding slide content.
        for line in (xml.title, *xml.body):
            if line and line not in seen:
                seen.add(line)
                body.append(line)
                if line != preferred.title:
                    recovered += 1
        note_lines = list(dict.fromkeys(
            line.strip()
            for value in (preferred.notes, xml.notes)
            for line in value.splitlines()
            if line.strip()
        ))
        merged.append(SlideText(number, title, tuple(body), "\n".join(note_lines)))

    if recovered:
        logger.info(
            "slide text supplemented from package XML slides=%d fields=%d",
            len(merged), recovered,
        )
    return merged or primary or fallback


def _preflight_zip_directory(path: Path) -> None:
    """Bound central-directory parsing before ``ZipFile`` allocates its index.

    The standard library reads the whole directory into memory before yielding
    ``ZipInfo`` objects. Its EOCD helper only reads the final 64 KiB, so inspect
    the advertised count/size first and reject pathological indexes early.
    """
    try:
        with path.open("rb") as source:
            end_record = zipfile._EndRecData(source)
    except OSError as exc:
        raise PresentationError("فایل ارائه قابل‌خواندن نیست.") from exc
    if end_record is None:
        return  # ``is_zipfile``/``ZipFile`` below will report malformed input.
    entries = int(end_record[zipfile._ECD_ENTRIES_TOTAL])
    directory_size = int(end_record[zipfile._ECD_SIZE])
    if entries > MAX_ZIP_ENTRIES:
        raise PresentationError("تعداد بخش‌های داخل فایل ارائه بیش از حد مجاز است.")
    if directory_size < 0 or directory_size > MAX_ZIP_CENTRAL_DIRECTORY_BYTES:
        raise PresentationError("فهرست داخلی فایل ارائه بیش از حد بزرگ است.")


def _directory_size(directory: Path) -> int:
    """Best-effort byte count of regular files currently under a work directory."""
    total = 0
    for item in directory.rglob("*"):
        try:
            if item.is_file():
                total += item.stat().st_size
        except OSError:
            continue
    return total


def _copy_zip_member_bounded(
    archive: zipfile.ZipFile,
    info: zipfile.ZipInfo,
    destination: Path,
    *,
    limit: int,
) -> int:
    """Copy a member with a streaming output cap and verify its declared size."""
    if info.file_size < 0 or info.file_size > limit:
        raise PresentationError("یکی از رسانه‌های ارائه از حد مجاز بزرگ‌تر است.")
    copied = 0
    stream_complete = False
    try:
        with archive.open(info) as source, destination.open("wb") as target:
            # Read at most one byte past the declared size to detect a forged
            # small file_size; the independent resource limit is checked too.
            source._left = min(limit, info.file_size) + 1
            while True:
                chunk = source.read(min(1024 * 1024, min(limit, info.file_size) - copied + 1))
                if not chunk:
                    break
                copied += len(chunk)
                if copied > info.file_size:
                    raise PresentationError("اندازهٔ واقعی رسانه از مقدار اعلام‌شده بیشتر است.")
                if copied > limit:
                    raise PresentationError("حجم بازشدهٔ رسانه از حد مجاز بیشتر است.")
                target.write(chunk)
            stream_complete = _zip_member_stream_complete(source, info)
    except PresentationError:
        destination.unlink(missing_ok=True)
        raise
    except (zipfile.BadZipFile, EOFError, OSError, RuntimeError) as exc:
        destination.unlink(missing_ok=True)
        raise PresentationError("یکی از رسانه‌های ارائه خراب است.") from exc
    if copied != info.file_size or not stream_complete:
        destination.unlink(missing_ok=True)
        raise PresentationError("اندازهٔ واقعی رسانه با ساختار فایل ارائه سازگار نیست.")
    return copied


def read_presentation(
    path: Path, media_dir: Path, settings: Settings
) -> PresentationContent:
    """Extract slide-linked media and slide text from a .pptx package (blocking)."""
    try:
        source_size = path.stat().st_size
    except OSError as exc:
        raise PresentationError("فایل ارائه قابل‌خواندن نیست.") from exc
    if source_size > min(settings.max_file_size, settings.presentation_max_temp_bytes):
        raise PresentationError("حجم فایل ارائه از سقف تعیین‌شده بیشتر است.")
    if not zipfile.is_zipfile(path):
        raise PresentationError(
            "این فایل یک ارائهٔ معتبر PowerPoint نیست یا آسیب دیده است."
        )
    _preflight_zip_directory(path)
    media_dir.mkdir(parents=True, exist_ok=True)
    clips: list[MediaClip] = []
    skipped: list[str] = []
    copied_paths: dict[str, Path] = {}
    copied_bytes = 0
    base_temp_bytes = _directory_size(path.parent)
    if base_temp_bytes > settings.presentation_max_temp_bytes:
        raise PresentationError("حجم فایل‌های موقت ارائه از سقف تعیین‌شده بیشتر است.")

    with zipfile.ZipFile(path) as archive:
        _validate_archive(archive, settings)
        names = set(archive.namelist())
        slide_parts = _slide_parts_in_order(archive)
        references: list[tuple[int | None, str | None, str]] = []
        referenced_parts: set[str] = set()
        slide_rels_by_part: dict[str, list[tuple[str, str, str]]] = {}
        for slide_number, slide_part in enumerate(slide_parts, start=1):
            all_slide_rels = _relationships(archive, slide_part)
            slide_rels_by_part[slide_part] = all_slide_rels
            slide_rels: list[tuple[str, str, str]] = []
            for rel_id, rel_type, target in all_slide_rels:
                if not target or rel_type not in MEDIA_REL_TYPES:
                    continue
                resolved = _resolve_target(slide_part, target)
                if resolved in names:
                    slide_rels.append((rel_id, MEDIA_REL_TYPES[rel_type], resolved))
                    referenced_parts.add(resolved)
            # Desktop PowerPoint often writes both an OOXML audio/video
            # relationship and a 2007 p14:media relationship for the same
            # embedded clip on a slide; drop the redundant "media" twin while
            # keeping explicit audio/video uses and standalone media links.
            primary_targets = {
                resolved for _rel_id, kind_tag, resolved in slide_rels if kind_tag != "media"
            }
            seen_media_targets: set[str] = set()
            for rel_id, kind_tag, resolved in slide_rels:
                if kind_tag == "media":
                    if resolved in primary_targets or resolved in seen_media_targets:
                        continue
                    seen_media_targets.add(resolved)
                references.append((slide_number, rel_id, resolved))

        unreferenced = sorted(
            (
                name for name in names
                if name.startswith("ppt/media/")
                and media_kind(name)
                and name not in referenced_parts
            ),
            key=natural_key,
        )
        if settings.presentation_include_unreferenced_media:
            references.extend((None, None, name) for name in unreferenced)
        elif unreferenced:
            skipped.append(
                f"رسانهٔ بدون پیوند مستقیم به اسلاید نادیده گرفته شد ({len(unreferenced)} فایل)"
            )

        for slide_number, rel_id, part_name in references:
            kind = media_kind(part_name)
            if kind is None:
                reason = f"{posixpath.basename(part_name)}: قالب رسانه پشتیبانی‌نشده"
                if reason not in skipped:
                    skipped.append(reason)
                continue
            if kind == "video" and not settings.presentation_include_video_audio:
                reason = f"{posixpath.basename(part_name)}: ویدیو نادیده گرفته شد"
                if reason not in skipped:
                    skipped.append(reason)
                continue
            if len(clips) >= settings.presentation_max_clips:
                if "تعداد استفاده‌های رسانه‌ای بیش از حد مجاز بود" not in skipped:
                    skipped.append("تعداد استفاده‌های رسانه‌ای بیش از حد مجاز بود")
                break

            try:
                info = archive.getinfo(part_name)
            except KeyError:
                continue
            if info.file_size > settings.max_file_size:
                reason = f"{posixpath.basename(part_name)}: حجم بیش از حد مجاز"
                if reason not in skipped:
                    skipped.append(reason)
                continue

            destination = copied_paths.get(part_name)
            if destination is None:
                if base_temp_bytes + copied_bytes + info.file_size > settings.presentation_max_temp_bytes:
                    raise PresentationError(
                        "حجم فایل‌های موقت ارائه از سقف تعیین‌شده بیشتر می‌شود."
                    )
                destination = media_dir / _safe_media_filename(
                    len(copied_paths) + 1, part_name
                )
                remaining_temp = settings.presentation_max_temp_bytes - base_temp_bytes - copied_bytes
                copied = _copy_zip_member_bounded(
                    archive,
                    info,
                    destination,
                    limit=min(settings.max_file_size, remaining_temp),
                )
                copied_paths[part_name] = destination
                copied_bytes += copied

            clips.append(
                MediaClip(
                    order=len(clips) + 1,
                    slide_number=slide_number,
                    part_name=part_name,
                    path=destination,
                    kind=kind,
                    relationship_id=rel_id,
                )
            )

        slides = (
            _extract_slide_text(
                path,
                archive=archive,
                slide_parts=slide_parts,
                slide_rels_by_part=slide_rels_by_part,
            )
            if settings.presentation_include_slide_text
            else []
        )
    if (
        settings.presentation_include_slide_text
        and slide_parts
        and not any(not slide.is_empty for slide in slides)
    ):
        logger.warning(
            "the deck has %d slide(s) but no readable slide text", len(slide_parts)
        )
        skipped.append(
            "متن قابل‌انتخاب روی اسلایدها پیدا نشد؛ جزوه فقط از صدای ارائه ساخته می‌شود"
        )
    return PresentationContent(tuple(slides), tuple(clips), tuple(skipped), len(slide_parts))


async def load_presentation(
    path: Path, media_dir: Path, settings: Settings
) -> PresentationContent:
    """Async wrapper so ZIP/XML work never blocks the Telethon event loop."""
    # Cancelling to_thread does not stop its worker. Join it before the caller
    # removes the job directory, otherwise extraction can recreate private files.
    worker = asyncio.create_task(asyncio.to_thread(read_presentation, path, media_dir, settings))
    try:
        return await asyncio.shield(worker)
    except asyncio.CancelledError:
        try:
            await worker
        except Exception:
            logger.warning("Presentation extraction failed during shutdown", exc_info=True)
        raise


async def prepare_audio(
    content: PresentationContent,
    workdir: Path,
    settings: Settings,
    *,
    skipped_out: list[str] | None = None,
) -> PreparedAudio | None:
    """Select, probe and prepare slide-linked clips for one deck-level STT call.

    Source duration excludes synthetic inter-clip silence and is the billed
    quantity. Output duration includes that silence and is the quantity checked
    against the presentation cap and used for timestamp-to-slide mapping.
    """
    skipped = list(content.skipped)

    def sync_skipped() -> None:
        if skipped_out is not None:
            skipped_out.clear()
            skipped_out.extend(skipped)

    sync_skipped()
    if not content.clips:
        return None

    # One worker process probes the whole deck. Duplicate paths are possible
    # when the same media part is referenced on more than one slide; probing it
    # once is enough, but each relationship remains a separate playback clip.
    unique_paths = list(dict.fromkeys(clip.path for clip in content.clips))
    probes = await probe_media_many(unique_paths, settings)
    usable: list[MediaClip] = []
    for clip in content.clips:
        info = probes.get(str(clip.path))
        if info is None:
            skipped.append(f"{clip.label}: بررسی فایل ناموفق بود")
            continue
        if not info.has_audio:
            skipped.append(f"{clip.label}: بدون شاخهٔ صوتی")
            continue
        if (
            info.duration is not None
            and info.duration < settings.presentation_min_clip_seconds
        ):
            skipped.append(f"{clip.label}: کوتاه‌تر از حد تعیین‌شده")
            continue
        usable.append(replace(clip, duration=info.duration))

    sync_skipped()
    if not usable:
        logger.info("No usable narration in the deck skipped=%s", len(skipped))
        return None

    measured_source = (
        sum(clip.duration for clip in usable)
        if all(clip.duration is not None for clip in usable)
        else None
    )
    estimated_output = (
        measured_source + settings.presentation_silence_seconds * (len(usable) - 1)
        if measured_source is not None else None
    )
    # A plainly oversized estimate can be rejected before decoding. The merge
    # report remains authoritative, since container probes can slightly
    # overestimate a video/audio stream's usable samples.
    if (
        estimated_output is not None
        and estimated_output > settings.presentation_max_total_duration + 1.0
    ):
        raise PresentationError(
            "مدت صدای ادغام‌شدهٔ این ارائه (با فاصله‌های بین کلیپ‌ها) از سقف ربات بیشتر است."
        )

    # The prepared audio exists only to be uploaded to an STT provider, so its
    # size budget is the tightest of the temp budget, MAX_FILE_SIZE and the
    # largest ceiling among the *configured* providers. Without this the merge
    # happily produced a WAV that every configured provider then rejected.
    stt_ceiling = largest_configured_upload(settings)
    upload_budget = (
        settings.max_file_size
        if stt_ceiling is None
        else min(settings.max_file_size, stt_ceiling)
    )

    if len(usable) == 1:
        clip = usable[0]
        info = probes[str(clip.path)]
        clip_size = await asyncio.to_thread(lambda p=clip.path: p.stat().st_size)
        if (
            clip.kind == "audio"
            and clip.duration is not None
            and clip_size <= upload_budget
            and not needs_transcode(clip.path.name, info)
        ):
            output_duration = clip.duration
            if output_duration > settings.presentation_max_total_duration:
                raise PresentationError(
                    "مدت صدای ارائه از سقف تعیین‌شدهٔ ربات بیشتر است."
                )
            if clip_size > settings.max_file_size:
                raise PresentationError("صدای آماده‌شده از سقف حجم ربات بزرگ‌تر است.")
            if await asyncio.to_thread(_directory_size, workdir) > settings.presentation_max_temp_bytes:
                raise PresentationError("حجم فایل‌های موقت ارائه از سقف تعیین‌شده بیشتر است.")
            return PreparedAudio(
                path=clip.path,
                clips=(clip,),
                source_duration=clip.duration,
                output_duration=output_duration,
                clip_spans=(ClipSpan(clip, 0.0, output_duration),),
                merged=False,
                skipped=tuple(skipped),
            )

    # If one duration is unavailable, choose a format for the configured
    # maximum rather than defaulting blindly to a potentially huge WAV.
    format_duration = (
        estimated_output
        if estimated_output is not None
        else settings.presentation_max_total_duration
    )
    available_temp_bytes = settings.presentation_max_temp_bytes - await asyncio.to_thread(
        _directory_size, workdir
    )
    if available_temp_bytes <= 0:
        raise PresentationError("برای آماده‌سازی صدای ارائه فضای موقت کافی نیست.")
    result = await merge_audio_tracks(
        [clip.path for clip in usable],
        workdir,
        settings,
        total_duration=format_duration,
        max_output_bytes=min(available_temp_bytes, upload_budget),
        stem="presentation-audio",
    )
    if result.output_duration > settings.presentation_max_total_duration + 1e-9:
        raise PresentationError(
            "مدت صدای ادغام‌شدهٔ این ارائه (با فاصله‌های بین کلیپ‌ها) از سقف ربات بیشتر است."
        )
    result_size = await asyncio.to_thread(lambda p=result.path: p.stat().st_size)
    if result_size > settings.max_file_size:
        raise PresentationError("صدای آماده‌شده از سقف حجم ربات بزرگ‌تر است.")
    if await asyncio.to_thread(_directory_size, workdir) > settings.presentation_max_temp_bytes:
        raise PresentationError("حجم فایل‌های موقت ارائه از سقف تعیین‌شده بیشتر است.")

    spans = tuple(
        ClipSpan(
            clip=usable[span.input_index],
            start_seconds=span.start_seconds,
            end_seconds=span.end_seconds,
        )
        for span in result.clip_spans
    )
    return PreparedAudio(
        path=result.path,
        clips=tuple(usable),
        source_duration=result.source_duration,
        output_duration=result.output_duration,
        clip_spans=spans,
        merged=True,
        skipped=tuple(skipped),
    )


def _match_clip_span(
    start: float,
    end: float,
    spans: list[ClipSpan],
    *,
    tolerance: float = ALIGNMENT_BOUNDARY_TOLERANCE_SECONDS,
) -> ClipSpan | None:
    """Return the unique clip span containing ``[start, end]``, with edge jitter tolerance.

    Exact containment wins first. When STT timestamps jitter slightly into the
    adjacent synthetic silence gap (up to ``tolerance`` seconds), a word is
    still attributed to its clip only if it overlaps that single clip, touches
    no other clip, and has at least half of its duration inside the clip.
    Words wholly inside a silence gap or crossing between clips return ``None``.
    """
    exact = [
        item for item in spans
        if start >= item.start_seconds and end <= item.end_seconds
    ]
    if len(exact) == 1:
        return exact[0]
    if len(exact) > 1 or tolerance <= 0:
        return None
    overlapping = [
        item for item in spans
        if start < item.end_seconds and end > item.start_seconds
    ]
    if len(overlapping) != 1:
        return None
    candidate = overlapping[0]
    leading_overhang = max(0.0, candidate.start_seconds - start)
    trailing_overhang = max(0.0, end - candidate.end_seconds)
    if leading_overhang > tolerance or trailing_overhang > tolerance:
        return None
    word_duration = end - start
    overlap_duration = min(end, candidate.end_seconds) - max(start, candidate.start_seconds)
    if word_duration > 0 and overlap_duration < word_duration * 0.5:
        return None
    return candidate


def align_narration_to_slides(
    raw: RawTranscript,
    prepared: PreparedAudio,
    slides: tuple[SlideText, ...] | list[SlideText],
) -> PresentationNarration:
    """Assign timed STT segments only to known slide-linked clip intervals.

    Providers that return plain text (or segments without timestamps) stay
    unaligned. Words in synthetic gaps, unreferenced clips, and segments without
    usable timestamps are kept in an explicit unassigned section, never guessed
    onto the nearest slide.
    """
    if not raw.words:
        plain = build_document(raw).plain()
        return PresentationNarration((), plain, False)

    slide_order = {slide.number: index for index, slide in enumerate(slides)}
    assigned: dict[int, list[Word]] = {}
    unassigned: list[Word] = []
    mapped_count = 0
    has_timing = False
    spans = sorted(prepared.clip_spans, key=lambda span: span.start_seconds)
    # When clip spans are disjoint and sorted (always true for merged audio),
    # narrow each word lookup to the O(1) adjacent window via binary search
    # instead of scanning every clip span linearly on every word.
    disjoint_sorted = len(spans) > 4 and all(
        spans[i].start_seconds <= spans[i].end_seconds <= spans[i + 1].start_seconds
        for i in range(len(spans) - 1)
    )
    span_starts = [span.start_seconds for span in spans] if disjoint_sorted else []

    for word in raw.words:
        if not word.text.strip():
            continue
        start, end = word.start, word.end
        timed = (
            start is not None and end is not None
            and math.isfinite(start) and math.isfinite(end)
            and start >= 0 and end >= start
        )
        if not timed:
            unassigned.append(word)
            continue
        has_timing = True
        assert start is not None and end is not None
        if disjoint_sorted:
            lo = max(0, bisect_right(span_starts, start) - 1)
            hi = min(len(spans), bisect_right(span_starts, end) + 1)
            candidate_spans = spans[lo:hi]
        else:
            candidate_spans = spans
        span = _match_clip_span(start, end, candidate_spans)
        slide_number = span.clip.slide_number if span is not None else None
        if slide_number is None:
            unassigned.append(word)
            continue
        assigned.setdefault(slide_number, []).append(word)
        mapped_count += 1

    if not has_timing or mapped_count == 0:
        plain = build_document(raw).plain()
        return PresentationNarration((), plain, False)

    slide_text: list[tuple[int, str]] = []
    for number in sorted(
        assigned, key=lambda value: slide_order.get(value, len(slide_order) + value)
    ):
        words = assigned[number]
        text = build_document(
            RawTranscript(
                words=words,
                text=" ".join(word.text for word in words),
                confidence=raw.confidence,
                engine=raw.engine,
            )
        ).plain()
        if text:
            slide_text.append((number, text))
    unassigned_text = ""
    if unassigned:
        unassigned_text = build_document(
            RawTranscript(
                words=unassigned,
                text=" ".join(word.text for word in unassigned),
                confidence=raw.confidence,
                engine=raw.engine,
            )
        ).plain()
    return PresentationNarration(tuple(slide_text), unassigned_text, bool(slide_text))


def slides_outline(slides: tuple[SlideText, ...] | list[SlideText], limit: int | None = None) -> str:
    """Render all slide material; an explicit limit is only for previews."""
    blocks = [slide.as_markdown() for slide in slides if not slide.is_empty]
    outline = "\n\n".join(blocks)
    if limit is not None and len(outline) > limit:
        outline = outline[:limit].rsplit("\n", 1)[0] + "\n\n(ادامهٔ متن اسلایدها کوتاه شد)"
    return outline
